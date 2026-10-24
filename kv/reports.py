# -*- coding: utf-8 -*-
"""Báo cáo cuối buổi: thống kê vi phạm theo khoảng thời gian và xuất Excel."""

import json
import os
import time


from .alerts import ALERT_FILE, ALERT_LOCK, ALERTS
from .evidence import evidence_root, watch_num
from .paths import CONFIG_PATH
from .privacy import images_blocked
from .runtime import ApiError, CFG


# ---------------------------------------------------------------- báo cáo cuối buổi
RANGES = {'2h': 'Hai giờ qua', 'today': 'Hôm nay', 'week': 'Tuần này', '7d': '7 ngày qua', 'month': 'Tháng này', '30d': '30 ngày qua', 'all': 'Tất cả'}
WEEKDAYS = ['Thứ hai', 'Thứ ba', 'Thứ tư', 'Thứ năm', 'Thứ sáu', 'Thứ bảy', 'Chủ nhật']
REPEAT_MIN_DAYS = 2      # học sinh vi phạm từ chừng này ngày khác nhau trở lên thì tính là tái phạm
CHRONIC_MIN_DAYS = 3     # trang/ứng dụng bị vào từ chừng này ngày khác nhau trở lên thì tính là lặp lại


def report_root():
    d = os.path.join(os.path.dirname(CONFIG_PATH), 'reports')
    os.makedirs(d, exist_ok=True)
    return d


def report_events(rng):
    import glob
    root = evidence_root()
    seen = {}
    for pth in [os.path.join(root, ALERT_FILE)] + sorted(glob.glob(os.path.join(root, 'alerts-*.old'))):
        try:
            with open(pth, 'r', encoding='utf-8') as f:
                for ln in f:
                    try:
                        e = json.loads(ln)
                        seen[(e['host'], e['file'])] = e
                    except (ValueError, KeyError, TypeError):
                        pass
        except OSError:
            pass
    with ALERT_LOCK:
        for e in ALERTS:
            seen[(e['host'], e['file'])] = e
    now = time.time()
    t = time.localtime()
    midnight = time.mktime((t.tm_year, t.tm_mon, t.tm_mday, 0, 0, 0, 0, 0, -1))
    lo = {'2h': now - 7200, 'today': midnight, 'week': midnight - t.tm_wday * 86400,
          '7d': now - 7 * 86400, 'month': time.mktime((t.tm_year, t.tm_mon, 1, 0, 0, 0, 0, 0, -1)),
          '30d': now - 30 * 86400, 'all': 0}[rng]
    return sorted((e for e in seen.values() if e.get('ts', 0) >= lo), key=lambda e: e['ts'])


def recent_text(rec, sep=' | '):
    """Nhật ký 5 phút của học sinh thành một dòng chữ: giờ ứng dụng [trang web] tiêu đề (thời gian), cũ -> mới."""
    out = []
    for x in rec or []:
        if not isinstance(x, dict):
            continue
        d = x.get('d') or 0
        out.append('%s %s%s%s (%s)' % (x.get('t', ''), x.get('p', ''), ' [%s]' % x['h'] if x.get('h') else '',
                                       ' ' + x['w'] if x.get('w') else '', '%ds' % d if d < 60 else '%d phút' % round(d / 60)))
    return sep.join(out)


def report_episodes(events, gap):
    """Gộp các ảnh liên tiếp của cùng máy + cùng trang/ứng dụng (cách nhau <= gap giây) thành một lần vi phạm."""
    out, last = [], {}
    for e in events:
        # tách theo lớp và học sinh: máy bỏ mở một trang qua giờ ra chơi thì ảnh của lớp sau không được gộp vào lần vi phạm của lớp trước
        k = (e['host'], e['type'], e['name'], e.get('cls', ''), e.get('student', ''))
        g = last.get(k)
        if g and e['ts'] - g['last_ts'] <= gap:
            g['shots'] += 1
            g['last_ts'] = e['ts']
            if e.get('img') and len(g['imgs']) < 3:
                g['imgs'].append(e['img'])
        else:
            g = {'host': e['host'], 'type': e['type'], 'name': e['name'], 'user': e.get('user', ''), 'window': e.get('window', ''),
                 'date': e['date'], 'time': e['time'], 'first_ts': e['ts'], 'last_ts': e['ts'], 'shots': 1,
                 'student': e.get('student', ''), 'cls': e.get('cls', ''), 'seat': e.get('seat', 0), 'recent': e.get('recent') or [], 'imgs': [e['img']] if e.get('img') else []}
            out.append(g)
            last[k] = g
    return out


def report_trend(eps):
    """Xu hướng theo thời gian từ các lần vi phạm: theo ngày, theo thứ trong tuần, học sinh tái phạm
    (vi phạm ở nhiều ngày khác nhau) và trang/ứng dụng bị vào lặp lại nhiều ngày."""
    from collections import Counter
    per_day, per_wd = Counter(), Counter()
    stu, item = {}, {}
    for g in eps:
        per_day[g['date']] += 1
        per_wd[time.localtime(g['first_ts']).tm_wday] += 1
        if g['student']:
            s = stu.setdefault((g['student'], g['cls']), {'days': set(), 'count': 0})
            s['days'].add(g['date'])
            s['count'] += 1
        it = item.setdefault((g['type'], g['name']), {'days': set(), 'count': 0, 'hosts': set()})
        it['days'].add(g['date'])
        it['count'] += 1
        it['hosts'].add(g['host'])
    repeat = sorted(({'student': k[0], 'cls': k[1], 'days': len(v['days']), 'count': v['count']}
                     for k, v in stu.items() if len(v['days']) >= REPEAT_MIN_DAYS),
                    key=lambda x: (-x['days'], -x['count'], x['student']))
    chronic = sorted(({'type': k[0], 'name': k[1], 'days': len(v['days']), 'count': v['count'], 'hosts': len(v['hosts'])}
                      for k, v in item.items() if len(v['days']) >= CHRONIC_MIN_DAYS),
                     key=lambda x: (-x['days'], -x['count'], x['name']))[:30]
    return {'by_day': [{'date': d, 'count': c} for d, c in sorted(per_day.items())],
            'by_weekday': [{'wd': i, 'label': WEEKDAYS[i], 'count': per_wd.get(i, 0)} for i in range(7)],
            'repeat_students': repeat[:50], 'chronic_items': chronic}


def report_build(rng, cfg=None):
    from collections import Counter
    if rng not in RANGES:
        raise ApiError('Khoảng thời gian không hợp lệ')
    cfg = cfg or CFG
    gap = watch_num(cfg, 'alert_group_sec', 60, 10, 3600)
    ev = report_events(rng)
    eps = report_episodes(ev, gap)
    hosts, items, hours = {}, {}, Counter()
    for g in eps:
        h = hosts.setdefault(g['host'], {'host': g['host'], 'count': 0, 'shots': 0, 'items': Counter(), 'last': ''})
        h['count'] += 1
        h['shots'] += g['shots']
        h['items'][(g['type'], g['name'])] += 1
        h['last'] = g['time']
        it = items.setdefault((g['type'], g['name']), {'type': g['type'], 'name': g['name'], 'count': 0, 'hosts': set()})
        it['count'] += 1
        it['hosts'].add(g['host'])
        hours[g['time'][:2]] += 1
    by_host = sorted(({'host': h['host'], 'count': h['count'], 'shots': h['shots'], 'last': h['last'],
                       'top': [{'type': k[0], 'name': k[1], 'count': c} for k, c in h['items'].most_common(3)]}
                      for h in hosts.values()), key=lambda x: (-x['count'], x['host']))
    by_item = sorted(({'type': i['type'], 'name': i['name'], 'count': i['count'], 'hosts': len(i['hosts'])} for i in items.values()),
                     key=lambda x: (-x['count'], x['name']))[:50]
    stu = {}
    for g in eps:
        key = (g['student'] or '(chưa biết học sinh)', g['cls'])
        r = stu.setdefault(key, {'student': key[0], 'cls': key[1], 'hosts': set(), 'count': 0, 'shots': 0, 'items': Counter(), 'last': ''})
        r['hosts'].add(g['host'])
        r['count'] += 1
        r['shots'] += g['shots']
        r['items'][(g['type'], g['name'])] += 1
        r['last'] = g['time']
    trend = report_trend(eps)
    by_student = sorted(({'student': r['student'], 'cls': r['cls'], 'hosts': sorted(r['hosts']), 'count': r['count'], 'shots': r['shots'],
                          'last': r['last'], 'top': [{'type': k[0], 'name': k[1], 'count': c} for k, c in r['items'].most_common(3)]}
                         for r in stu.values()), key=lambda x: (-x['count'], x['student']))
    return {'range': rng, 'label': RANGES[rng], 'gap': gap, 'by_student': by_student, 'generated': time.strftime('%Y-%m-%d %H:%M:%S'),
            'episodes': len(eps), 'shots': len(ev), 'machines': len(hosts),
            'first': ev[0]['time'] if ev else '', 'last': ev[-1]['time'] if ev else '',
            'by_host': by_host, 'by_item': by_item, 'by_day': trend['by_day'], 'by_weekday': trend['by_weekday'],
            'repeat_students': trend['repeat_students'], 'chronic_items': trend['chronic_items'],
            'by_hour': [{'hour': h, 'count': c} for h, c in sorted(hours.items())],
            # trang nhiều máy cùng vào: có thể là trang cần cho bài học, đáng cân nhắc thêm vào AllowPage
            'suggest': [i['name'] for i in by_item if i['type'] == 'site' and i['hosts'] >= 3],
            '_episodes': eps, '_events': ev}


def api_report(q):
    rep = report_build(q.get('range') or 'today')
    rep.pop('_episodes')
    rep.pop('_events')
    return rep


def report_write_xlsx(rep, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill  # nạp muộn: openpyxl khá nặng, chỉ cần khi đọc/ghi Excel
    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='2563EB')

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(list(r))
            for c in ws[ws.max_row]:  # chuỗi bắt đầu bằng = + - @ vẫn là chữ, không bị Excel hiểu thành công thức
                if isinstance(c.value, str):
                    c.data_type = 's'
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = 'A2'
    kind = lambda t: {'site': 'Trang web', 'ai': 'Công cụ AI'}.get(t, 'Ứng dụng')  # noqa
    ws = wb.active
    ws.title = 'Tổng quan'
    sheet(ws, ['Mục', 'Giá trị'], [
        ('Khoảng thời gian', rep['label']), ('Lập lúc', rep['generated']),
        ('Số lần vi phạm', rep['episodes']), ('Số ảnh chụp', rep['shots']), ('Số máy có vi phạm', rep['machines']),
        ('Vi phạm đầu tiên', rep['first']), ('Vi phạm cuối cùng', rep['last']),
        ('Ghi chú', 'Một lần vi phạm = các ảnh liên tiếp của cùng máy và cùng trang/ứng dụng, cách nhau không quá %d giây.' % rep['gap']),
        ('Gợi ý thêm vào AllowPage', ', '.join(rep['suggest']) or '(không có)'),
    ], [28, 90])
    sheet(wb.create_sheet('Theo máy'), ['Máy', 'Số lần vi phạm', 'Số ảnh', 'Lần cuối', 'Hay vào nhất'],
          [(h['host'], h['count'], h['shots'], h['last'], ', '.join('%s (%d)' % (t['name'], t['count']) for t in h['top'])) for h in rep['by_host']],
          [18, 16, 10, 12, 70])
    sheet(wb.create_sheet('Theo học sinh'), ['Học sinh', 'Lớp', 'Máy', 'Số lần vi phạm', 'Số ảnh', 'Lần cuối', 'Hay vào nhất'],
          [(x['student'], x['cls'], ', '.join(x['hosts']), x['count'], x['shots'], x['last'], ', '.join('%s (%d)' % (t['name'], t['count']) for t in x['top'])) for x in rep['by_student']],
          [28, 10, 20, 16, 10, 12, 70])
    sheet(wb.create_sheet('Theo trang-ứng dụng'), ['Nội dung', 'Loại', 'Số lần vi phạm', 'Số máy'],
          [(i['name'], kind(i['type']), i['count'], i['hosts']) for i in rep['by_item']], [44, 14, 16, 10])
    sheet(wb.create_sheet('Theo ngày'), ['Ngày', 'Số lần vi phạm'], [(d['date'], d['count']) for d in rep.get('by_day', [])], [14, 16])
    sheet(wb.create_sheet('Tái phạm'), ['Học sinh', 'Lớp', 'Số ngày vi phạm', 'Số lần vi phạm'],
          [(x['student'], x['cls'], x['days'], x['count']) for x in rep['repeat_students']], [28, 10, 18, 16])
    sheet(wb.create_sheet('Lặp lại nhiều ngày'), ['Nội dung', 'Loại', 'Số ngày', 'Số lần', 'Số máy'],
          [(x['name'], kind(x['type']), x['days'], x['count'], x['hosts']) for x in rep.get('chronic_items', [])], [44, 14, 10, 10, 10])
    sheet(wb.create_sheet('Chi tiết'), ['Ngày', 'Giờ', 'Học sinh', 'Lớp', 'Máy', 'Tài khoản', 'Loại', 'Nội dung', 'Tiêu đề cửa sổ', 'Ảnh', '5 phút trước đó (cũ → mới)'],
          [(e['date'], e['time'], e.get('student', ''), e.get('cls', ''), e['host'], e.get('user', ''), kind(e['type']), e['name'], e.get('window', ''), e.get('img', ''), recent_text(e.get('recent'))) for e in rep['_events']],
          [12, 10, 26, 8, 16, 16, 12, 36, 60, 60, 120])
    wb.save(path)


def _img_data_uri(rel, max_side=560):
    """Ảnh vi phạm thành data URI (thu nhỏ nếu có Pillow) để file HTML tự chứa đủ, gửi đi vẫn xem được."""
    import base64
    import io
    root = os.path.realpath(evidence_root())
    p = os.path.realpath(os.path.join(root, *str(rel).split('/')))
    if not p.startswith(root + os.sep) or not os.path.isfile(p):
        return ''
    try:
        raw = None
        try:
            from PIL import Image
            im = Image.open(p)
            im.thumbnail((max_side, max_side))
            buf = io.BytesIO()
            im.convert('RGB').save(buf, 'JPEG', quality=70)
            raw = buf.getvalue()
        except Exception:  # noqa
            if os.path.getsize(p) <= 450 * 1024:
                with open(p, 'rb') as f:
                    raw = f.read()
        return 'data:image/jpeg;base64,' + base64.b64encode(raw).decode('ascii') if raw else ''
    except OSError:
        return ''


HTML_MAX_IMAGES = 120   # tối đa số ảnh nhúng vào báo cáo in (giữ file nhẹ, đủ để nộp)


def report_write_html(rep, path, info=None):
    """Báo cáo để in hoặc lưu PDF (Ctrl+P): tổng quan, theo học sinh, từng lần vi phạm kèm ảnh. Mọi chuỗi đều được escape."""
    from html import escape as E
    info = info or {}
    kind = lambda t: {'site': 'Trang web', 'ai': 'Công cụ AI'}.get(t, 'Ứng dụng')  # noqa
    stu_rows = ''.join('<tr><td><b>%s</b></td><td>%s</td><td>%s</td><td class="n">%d</td><td>%s</td><td>%s</td></tr>' % (
        E(x['student']), E(x['cls']), E(', '.join(x['hosts'])), x['count'], E(x['last']),
        E(', '.join('%s (%d)' % (t['name'], t['count']) for t in x['top']))) for x in rep['by_student'])
    item_rows = ''.join('<tr><td>%s</td><td>%s</td><td class="n">%d</td><td class="n">%d</td></tr>' % (E(i['name']), kind(i['type']), i['count'], i['hosts'])
                        for i in rep['by_item'][:20])
    left, cards = HTML_MAX_IMAGES, []
    for g in sorted(rep['_episodes'], key=lambda g: (g['student'] or '~', g['first_ts'])):
        imgs = []
        for rel in g.get('imgs', []):
            if left <= 0:
                break
            uri = _img_data_uri(rel)
            if uri:
                imgs.append('<img src="%s" alt="">' % uri)
                left -= 1
        who = g['student'] or 'Chưa biết học sinh'
        cards.append('<section class="ep"><div class="h"><b>%s</b> <span>%s · máy %s%s</span></div><div class="m">%s %s · %s: <b>%s</b>%s</div>%s</section>' % (
            E(who), E(g['cls'] or ''), E(g['host']), (' · chỗ %d' % g['seat']) if g.get('seat') else '',
            E(g['date']), E(g['time']), kind(g['type']), E(g['name']),
            ('<div class="w">%s</div>' % E(g['window'])) if g.get('window') else '',
            (('<details class="w"><summary>5 phút trước đó</summary>%s</details>' % E(recent_text(g.get('recent'), ' › '))) if g.get('recent') else '') +
            (('<div class="im">%s</div>' % ''.join(imgs)) if imgs else '')))
    meta = ' · '.join(x for x in [info.get('cls') and 'Lớp %s' % info['cls'], info.get('subject'), info.get('teacher')] if x)
    html = """<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>Báo cáo vi phạm - %(label)s</title><style>
body{font:14px/1.45 "Segoe UI",Arial,sans-serif;color:#1b2540;margin:24px auto;max-width:900px;padding:0 16px}
h1{font-size:22px;margin:0 0 4px}h2{font-size:16px;margin:26px 0 8px;border-bottom:2px solid #2563eb;padding-bottom:4px}
.sub{color:#5b6786;margin:0 0 14px}.stats{display:flex;gap:10px;flex-wrap:wrap}.stats div{border:1px solid #d5dcec;border-radius:10px;padding:8px 14px}
.stats b{display:block;font-size:20px}table{border-collapse:collapse;width:100%%}th,td{border:1px solid #d5dcec;padding:5px 8px;text-align:left;vertical-align:top}
th{background:#eef3ff}td.n{text-align:right}.ep{border:1px solid #d5dcec;border-radius:10px;padding:8px 12px;margin:10px 0;break-inside:avoid}
.ep .h span{color:#5b6786}.ep .m{margin-top:2px}.ep .w{color:#5b6786;font-size:12px}.im{display:flex;gap:6px;margin-top:6px;flex-wrap:wrap}.im img{width:31%%;min-width:200px;border:1px solid #d5dcec;border-radius:6px}
.bar{position:sticky;top:0;background:#fff;padding:8px 0;text-align:right}.bar button{padding:6px 14px;border-radius:8px;border:1px solid #2563eb;background:#2563eb;color:#fff;cursor:pointer}
@media print{.bar{display:none}body{margin:0;max-width:none}@page{size:A4;margin:14mm}}
</style></head><body><div class="bar"><button onclick="window.print()">In / Lưu thành PDF</button></div>
<h1>Báo cáo vi phạm: %(label)s</h1><p class="sub">%(meta)s%(sep)sLập lúc %(gen)s</p>
<div class="stats"><div><b>%(ep)d</b>lần vi phạm</div><div><b>%(shots)d</b>ảnh chụp</div><div><b>%(mach)d</b>máy có vi phạm</div><div><b>%(nstu)d</b>học sinh</div></div>
<h2>Theo học sinh</h2><table><tr><th>Học sinh</th><th>Lớp</th><th>Máy</th><th>Số lần</th><th>Lần cuối</th><th>Hay vào nhất</th></tr>%(stu)s</table>
<h2>Trang web / ứng dụng vi phạm nhiều nhất</h2><table><tr><th>Nội dung</th><th>Loại</th><th>Số lần</th><th>Số máy</th></tr>%(items)s</table>
<h2>Xu hướng theo ngày</h2><table><tr><th>Ngày</th><th>Số lần vi phạm</th></tr>%(days)s</table>
<h2>Học sinh tái phạm (vi phạm ở nhiều ngày)</h2><table><tr><th>Học sinh</th><th>Lớp</th><th>Số ngày</th><th>Số lần</th></tr>%(repeat)s</table>
<h2>Từng lần vi phạm</h2>%(cards)s<p class="sub">Một lần vi phạm gồm các ảnh liên tiếp của cùng máy và cùng trang hoặc ứng dụng, cách nhau không quá %(gap)d giây. Tối đa %(maxi)d ảnh được nhúng. Tên học sinh lấy theo sơ đồ chỗ ngồi tại thời điểm vi phạm.</p>
</body></html>""" % {'label': E(rep['label']), 'meta': E(meta), 'sep': ' · ' if meta else '', 'gen': E(rep['generated']), 'ep': rep['episodes'], 'shots': rep['shots'],
                    'mach': rep['machines'], 'nstu': len([x for x in rep['by_student'] if x['student'] != '(chưa biết học sinh)']), 'stu': stu_rows or '<tr><td colspan="6">Không có vi phạm</td></tr>',
                    'items': item_rows or '<tr><td colspan="4">Không có vi phạm</td></tr>', 'cards': ''.join(cards) or '<p>Không có vi phạm.</p>',
                    'days': ''.join('<tr><td>%s</td><td class="n">%d</td></tr>' % (E(d['date']), d['count']) for d in rep.get('by_day', [])) or '<tr><td colspan="2">Không có vi phạm</td></tr>',
                    'repeat': ''.join('<tr><td><b>%s</b></td><td>%s</td><td class="n">%d</td><td class="n">%d</td></tr>' % (E(x['student']), E(x['cls']), x['days'], x['count']) for x in rep.get('repeat_students', [])) or '<tr><td colspan="4">Không có học sinh nào vi phạm từ %d ngày khác nhau trở lên</td></tr>' % REPEAT_MIN_DAYS, 'gap': rep['gap'], 'maxi': HTML_MAX_IMAGES}
    with open(path, 'w', encoding='utf-8') as f:
        f.write(html)


def api_report_export(b):
    rep = report_build(b.get('range') or 'today')
    fmt = 'html' if b.get('format') == 'html' else 'xlsx'
    if fmt == 'html' and images_blocked():
        raise ApiError('Báo cáo in có ảnh học sinh. Hãy nhập mã xem ảnh ở trang Vi phạm trước.')
    path = os.path.join(report_root(), 'BaoCao-%s-%s.%s' % (rep['range'], time.strftime('%Y%m%d-%H%M%S'), fmt))
    if fmt == 'html':
        from .school import school_now
        cur = (school_now()['current'] or {})
        from .school import SCHOOL, SCHOOL_LOCK
        with SCHOOL_LOCK:
            info = {'cls': cur.get('cls', ''), 'subject': SCHOOL.get('subject', ''), 'teacher': SCHOOL.get('teacher', '')}
        report_write_html(rep, path, info)
        if os.name == 'nt':
            try:
                os.startfile(path)  # noqa
            except OSError:
                pass
    else:
        report_write_xlsx(rep, path)
    return {'path': path, 'format': fmt}
