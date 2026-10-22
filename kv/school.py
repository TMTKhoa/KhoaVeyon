# -*- coding: utf-8 -*-
"""Dữ liệu lớp học: danh sách lớp, thời khóa biểu, tiết hiện tại, kiểm tra sơ đồ phòng máy."""

import datetime
import json
import os
import re
import threading
import time


from .paths import CONFIG_PATH
from .runtime import ApiError, log


SCHOOL_LOCK = threading.RLock()
SCHOOL_GRACE_MIN = 5       # vào lớp sớm tối đa chừng này phút vẫn tính là tiết đó
SCHOOL_MANUAL_MIN = 45     # "đặt làm lớp hiện tại" có hiệu lực chừng này phút (khi đang bật tự động theo lịch)
SCHOOL_DAYS = ['THỨ HAI', 'THỨ BA', 'THỨ TƯ', 'THỨ NĂM', 'THỨ SÁU', 'THỨ BẢY', 'CHỦ NHẬT']
SCHOOL_DAY_SHORT = ['T2', 'T3', 'T4', 'T5', 'T6', 'T7', 'CN']
SCHOOL = {'classes': {}, 'layouts': {}, 'slots': [], 'from': '', 'auto': True, 'override': None, 'manual': '',
          'classes_file': '', 'classes_at': '', 'timetable_file': '', 'timetable_at': '', 'teacher': '', 'subject': '',
          'bells': {}, 'wall_presets': {}, 'wall_follow': None, 'lock_style': None, 'hand_enabled': True, 'act_enabled': True, 'attendance': {}, 'attend_notify': True, 'auto_announce': False, 'kq_game_absent': True, 'automation': None}


def school_path():
    return os.path.join(os.path.dirname(CONFIG_PATH), 'school.json')


def school_load():
    try:
        with open(school_path(), 'r', encoding='utf-8') as f:
            d = json.load(f)
        if isinstance(d, dict):
            with SCHOOL_LOCK:
                SCHOOL.update({k: d[k] for k in SCHOOL if k in d})
    except (OSError, ValueError):
        pass


def school_save():
    with SCHOOL_LOCK:
        tmp = school_path() + '.tmp'
        try:
            os.makedirs(os.path.dirname(school_path()), exist_ok=True)
            with open(tmp, 'w', encoding='utf-8') as f:
                json.dump(SCHOOL, f, ensure_ascii=False, indent=1)
            os.replace(tmp, school_path())
        except OSError as e:
            log('Không lưu được lớp và lịch dạy: %s' % e, 'err')


def _cls_norm(v):
    return re.sub(r'\s+', '', str(v or '')).upper()


def _hhmm(v):
    """Giờ dạng '7:30', '07:30', datetime.time hoặc số của Excel -> 'HH:MM' (không hiểu thì trả về '')."""
    if isinstance(v, datetime.datetime):
        v = v.time()
    if isinstance(v, datetime.time):
        return '%02d:%02d' % (v.hour, v.minute)
    if isinstance(v, (int, float)) and 0 <= v < 1:
        m = int(round(v * 1440))
        return '%02d:%02d' % (m // 60, m % 60)
    m = re.match(r'^\s*(\d{1,2})\s*[:hH.]\s*(\d{2})', str(v or ''))
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return '%02d:%02d' % (int(m.group(1)), int(m.group(2)))
    return ''


def _date_iso(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    m = re.match(r'^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})', str(v or ''))
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).strftime('%Y-%m-%d')
        except ValueError:
            return ''
    return ''


def parse_classes_xlsx(path):
    """File điểm danh của trường: mỗi lớp một sheet, tên lớp ở ô E2, họ tên học sinh ở cột C (từ dưới dòng tiêu đề 'HỌ & TÊN')."""
    from openpyxl import load_workbook  # nạp muộn: openpyxl khá nặng, chỉ cần khi đọc/ghi Excel
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=75, max_col=5, values_only=True))
            if len(rows) < 3:
                continue
            row2 = rows[1]
            name = _cls_norm(row2[4] if len(row2) > 4 else '')
            if not name or 'Lớp' not in str(row2[3] if len(row2) > 3 else ''):
                continue
            head = None
            for i, r in enumerate(rows[:12]):
                if len(r) > 2 and 'HỌ' in str(r[2] or '').upper():
                    head = i
                    break
            if head is None:
                continue
            names = []
            for r in rows[head + 1:]:
                v = r[2] if len(r) > 2 else None
                if isinstance(v, str) and v.strip():
                    names.append(re.sub(r'\s+', ' ', v).strip())
            if names:
                out[name] = names
    finally:
        wb.close()
    if not out:
        raise ApiError('Không tìm thấy lớp nào. File cần có mỗi lớp một sheet, tên lớp ở ô E2 và họ tên học sinh ở cột C.')
    return out


def parse_classes_json(path):
    """File JSON danh sách lớp xuất từ KhoaTool (Thống kê → "Xuất cho KhoaVeyon"):
    {"app": "KhoaTool", "format": 1, "classes": [{"cls": "6A1", "students": ["Nguyễn An", ...]}]}.
    Cũng nhận dạng rút gọn {"6A1": ["Nguyễn An", ...]}. Trả về {lớp: [học sinh]}."""
    with open(path, 'r', encoding='utf-8-sig') as f:
        try:
            d = json.load(f)
        except ValueError:
            raise ApiError('File JSON không hợp lệ')
    pairs = []
    if isinstance(d, dict) and isinstance(d.get('classes'), list):
        pairs = [(c.get('cls'), c.get('students')) for c in d['classes'] if isinstance(c, dict)]
    elif isinstance(d, dict):
        pairs = list(d.items())
    out = {}
    for cls, names in pairs:
        cn = _cls_norm(cls)[:20]
        if not cn or not isinstance(names, list):
            continue
        seen, lst = set(), []
        for n in names:
            n = re.sub(r'\s+', ' ', str(n or '')).strip()[:80]
            if n and n.lower() not in seen:
                seen.add(n.lower())
                lst.append(n)
        if lst:
            out.setdefault(cn, [])
            out[cn] += [n for n in lst if n not in out[cn]]
    if not out:
        raise ApiError('Không tìm thấy lớp nào trong file JSON (cần dạng {"classes": [{"cls": "6A1", "students": [...]}]})')
    return out


def _fold(v):
    """Chữ hoa, bỏ dấu tiếng Việt và khoảng trắng thừa: 'Thứ  Hai' -> 'THU HAI'."""
    import unicodedata
    t = str(v or '').replace('Đ', 'D').replace('đ', 'd')
    t = ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', t).strip().upper()


_DAY_WORDS = {'HAI': 0, '2': 0, 'BA': 1, '3': 1, 'TU': 2, '4': 2, 'NAM': 3, '5': 3, 'SAU': 4, '6': 4, 'BAY': 5, '7': 5}


def _day_of(v):
    """Ô tiêu đề cột -> số thứ (0 = thứ hai) hoặc None. Nhận: THỨ HAI, Thu 2, T2, THỨ 2 (19/10), Monday, CN, Chủ nhật..."""
    t = _fold(v)
    if not t:
        return None
    if re.match(r'^(CHU NHAT|CN|SUNDAY|SUN)\b', t):
        return 6
    m = re.match(r'^(?:THU|T)\s*(HAI|BA|TU|NAM|SAU|BAY|[2-7])\b', t)
    if m:
        return _DAY_WORDS[m.group(1)]
    for k, en in enumerate(('MONDAY', 'TUESDAY', 'WEDNESDAY', 'THURSDAY', 'FRIDAY', 'SATURDAY')):
        if t.startswith(en[:3]) and (len(t) == 3 or t.startswith(en)):
            return k
    return None


# Giờ tiết mặc định khi file thời khóa biểu chỉ ghi số tiết, không ghi giờ (sáng 7:00, chiều 13:30, mỗi tiết 45 phút)
DEFAULT_BELLS = {1: ('07:00', '07:45'), 2: ('07:50', '08:35'), 3: ('08:55', '09:40'), 4: ('09:45', '10:30'), 5: ('10:35', '11:20'),
                 6: ('13:30', '14:15'), 7: ('14:20', '15:05'), 8: ('15:25', '16:10'), 9: ('16:15', '17:00'), 10: ('17:05', '17:50')}


def school_bells():
    """Giờ 10 tiết đang dùng: giờ giáo viên đã đặt (SCHOOL['bells']), tiết nào chưa đặt thì lấy giờ mặc định."""
    out = dict(DEFAULT_BELLS)
    with SCHOOL_LOCK:
        for k, v in (SCHOOL.get('bells') or {}).items():
            try:
                n = int(k)
                if 1 <= n <= 10 and _hhmm(v[0]) and _hhmm(v[1]):
                    out[n] = (_hhmm(v[0]), _hhmm(v[1]))
            except (ValueError, TypeError, IndexError):
                pass
    return out


def _times_in_row(cells):
    """Các giờ có trong một dãy ô, kể cả khi hai giờ nằm chung một ô ('7:00 - 7:45', '07h00-07h45')."""
    out = []
    for v in cells:
        t = '' if isinstance(v, str) else _hhmm(v)
        if t:
            out.append(t)
            continue
        out += ['%02d:%02d' % (int(a), int(b)) for a, b in re.findall(r'(\d{1,2})\s*[:hH]\s*(\d{2})', str(v or '')) if int(a) < 24 and int(b) < 60]
    return out


def _period_of(cells):
    for v in cells:
        m = re.match(r'^\s*(?:TIET|T)?\s*(\d{1,2})\s*$', _fold(v))
        if m and 1 <= int(m.group(1)) <= 10:
            return int(m.group(1))
    return 0


def _cls_from_cell(v):
    """Ô tiết học -> tên lớp. 'Lớp 10A1 (P.201)' -> '10A1'; ô trống hoặc chỉ ghi '-' -> ''."""
    raw = str(v or '').strip()
    if not raw or raw in ('-', '--', '/', 'x', 'X'):
        return ''
    m = re.search(r'(?<![0-9A-Za-zÀ-ỹ])(\d{1,2}\s?[A-Za-zĐđ]{1,3}\s?\d{0,2}(?:\.\d{1,2})?)(?![0-9A-Za-zÀ-ỹ])', raw)
    return _cls_norm(m.group(1)) if m else _cls_norm(raw)


def parse_timetable_xlsx(path):
    """Thời khóa biểu dạng lưới: hàng tiêu đề có các thứ (THỨ HAI, Thu 2, T2...), bên trái là giờ (hai cột giờ bắt đầu/kết thúc,
    hoặc một ô '7:00 - 7:45', hoặc số tiết), các ô còn lại là tên lớp. Xét mọi sheet, lấy sheet có tiêu đề các thứ."""
    from openpyxl import load_workbook  # nạp muộn: openpyxl khá nặng, chỉ cần khi đọc/ghi Excel
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        sheets = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()
    seen_titles = []
    for title, grid in sheets:
        hdr = cols = None
        for i, r in enumerate(grid[:60]):
            found = {}
            for j, v in enumerate(r):
                d = _day_of(v)
                if d is not None and d not in found.values():
                    found[j] = d
            if len(found) >= 2:
                hdr, cols = i, found
                break
        if hdr is None:
            seen_titles.append(title)
            continue
        first = min(cols)
        slots, frm, teacher, subject, used_default = [], '', '', '', False
        for r in grid[:hdr]:
            for j, v in enumerate(r):
                t = _fold(v)
                nxt = next((x for x in r[j + 1:] if x not in (None, '')), None)
                if t in ('BAT DAU TU', 'AP DUNG TU', 'TU NGAY'):
                    frm = _date_iso(nxt)
                elif t in ('GIAO VIEN', 'GV') and nxt:
                    teacher = str(nxt).strip()
                elif t in ('MON', 'MON HOC') and nxt:
                    subject = str(nxt).strip()
        for r in grid[hdr + 1:]:
            left = r[:first]
            times = _times_in_row(left)
            if len(times) >= 2:
                st, en = times[0], times[1]
            else:
                per = _period_of(left)
                bells = school_bells()
                if per and per in bells:
                    st, en = bells[per]
                    used_default = True
                else:
                    continue
            for j, day in cols.items():
                c = _cls_from_cell(r[j] if j < len(r) else '')
                if c:
                    slots.append({'day': day, 'start': st, 'end': en, 'cls': c})
        if not slots:
            raise ApiError('Thấy hàng tiêu đề các thứ ở sheet "%s" nhưng không có tiết nào có tên lớp. Mỗi hàng tiết cần có giờ (7:00 và 7:45, hoặc "7:00 - 7:45", hoặc số tiết 1-10) ở bên trái các cột thứ.' % title)
        slots.sort(key=lambda s: (s['day'], s['start']))
        return {'slots': slots, 'from': frm, 'teacher': teacher, 'subject': subject, 'default_bells': used_default}
    raise ApiError('Không thấy hàng tiêu đề các thứ (THỨ HAI, Thứ 2, T2...) ở sheet nào (đã xem: %s). Hàng tiêu đề cần ít nhất 2 cột thứ.' % ', '.join(seen_titles[:6]))


def _mins(hhmm):
    return int(hhmm[:2]) * 60 + int(hhmm[3:5])


def school_active_slot(dt):
    """Tiết đang diễn ra tại thời điểm dt (kể cả vào sớm SCHOOL_GRACE_MIN phút), theo thời khóa biểu."""
    with SCHOOL_LOCK:
        frm, slots = SCHOOL.get('from') or '', list(SCHOOL.get('slots') or [])
    if frm and dt.strftime('%Y-%m-%d') < frm:
        return None
    m = dt.hour * 60 + dt.minute
    for s in slots:
        if s['day'] == dt.weekday() and _mins(s['start']) - SCHOOL_GRACE_MIN <= m < _mins(s['end']):
            return s
    return None


def school_next_slot(dt):
    with SCHOOL_LOCK:
        frm, slots = SCHOOL.get('from') or '', list(SCHOOL.get('slots') or [])
    for add in range(0, 8):
        d = dt + datetime.timedelta(days=add)
        if frm and d.strftime('%Y-%m-%d') < frm:
            continue
        m = dt.hour * 60 + dt.minute if add == 0 else -1
        for s in slots:
            if s['day'] == d.weekday() and _mins(s['start']) - (SCHOOL_GRACE_MIN if add == 0 else 0) > m:
                return dict(s, date=d.strftime('%Y-%m-%d'))
    return None


def school_now(ts=None):
    """Lớp hiện tại và lớp kế tiếp. source: schedule (theo lịch), manual (giáo viên đặt tay), none."""
    ts = time.time() if ts is None else ts
    dt = datetime.datetime.fromtimestamp(ts)
    with SCHOOL_LOCK:
        auto, ov, manual = bool(SCHOOL.get('auto', True)), SCHOOL.get('override'), SCHOOL.get('manual') or ''
    cur, source = None, 'none'
    if not auto:
        if manual:
            cur, source = {'cls': manual}, 'manual'
    elif ov and ov.get('since', 0) <= ts < ov.get('until', 0):
        cur, source = {'cls': ov['cls'], 'until': ov['until']}, 'manual'
    else:
        s = school_active_slot(dt)
        if s:
            cur, source = dict(s), 'schedule'
    nxt = school_next_slot(dt) if auto else None
    if nxt and cur and nxt.get('cls') == cur.get('cls') and nxt.get('start') == cur.get('start'):
        nxt = None
    return {'current': cur, 'next': nxt, 'source': source}


def school_class_at(ts):
    """Tên lớp tại thời điểm ts, dùng để ghi vào mỗi vi phạm (rỗng nếu không biết)."""
    c = school_now(ts)['current']
    return (c or {}).get('cls', '')


LAYOUT_DEFAULT = {'cols': 4, 'per': 10, 'from': 'teacher', 'first': 'left', 'pattern': '', 'seats': {}, 'broken': [], 'gv': 'center', 'door': 'back', 'flipv': False, 'fliph': False}


def layout_clean(b, old_seats=None, old_broken=None):
    """Kiểm tra sơ đồ phòng máy do giao diện gửi lên: số dãy, máy mỗi dãy, đánh số từ bàn giáo viên hay từ cửa, dãy 1 bên nào."""
    try:
        cols, per = int(b.get('cols')), int(b.get('per'))
    except (TypeError, ValueError):
        raise ApiError('Số dãy và số máy mỗi dãy phải là số')
    if not (1 <= cols <= 12 and 1 <= per <= 20):
        raise ApiError('Số dãy từ 1 đến 12 và số máy mỗi dãy từ 1 đến 20')
    frm = b.get('from') if b.get('from') in ('teacher', 'door') else 'teacher'
    first = b.get('first') if b.get('first') in ('left', 'right') else 'left'
    pat = str(b.get('pattern') or '').strip()[:40]
    if pat and '{n}' not in pat:
        raise ApiError('Tên máy mẫu phải có {n} thay cho số máy, ví dụ MAY{n}')
    seats = b.get('seats') if isinstance(b.get('seats'), dict) else (old_seats or {})
    clean = {}
    for k, v in seats.items():
        if str(k).isdigit() and 1 <= int(k) <= cols * per and isinstance(v, str) and v.strip():
            clean[str(int(k))] = v.strip()[:80]
    gv = b.get('gv') if b.get('gv') in ('left', 'center', 'right') else 'center'
    door = b.get('door') if b.get('door') in ('back', 'tl', 'tr') else 'back'
    broken = b.get('broken') if isinstance(b.get('broken'), list) else (old_broken or [])
    bk = sorted({int(x) for x in broken if str(x).isdigit() and 1 <= int(x) <= cols * per})
    for k in bk:
        clean.pop(str(k), None)  # máy hỏng thì không có người ngồi
    return {'cols': cols, 'per': per, 'from': frm, 'first': first, 'pattern': pat, 'seats': clean,
            'broken': bk, 'gv': gv, 'door': door, 'flipv': b.get('flipv') is True, 'fliph': b.get('fliph') is True}
