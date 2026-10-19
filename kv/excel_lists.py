# -*- coding: utf-8 -*-
"""Bốn danh sách Excel (KeepAppList, AllowPage, BlockAppList, BlockAIList): mẫu có sẵn, đọc, ghi, kiểm tra từng dòng."""

import os
import re


# ---------------------------------------------------------------- danh sách Excel
AI_PRESETS = {
    'ChatGPT, OpenAI': [
        ('chatgpt.com', 'ChatGPT'), ('openai.com', 'OpenAI, gồm chat.openai.com và API'), ('sora.com', 'Sora tạo video'),
    ],
    'Google Gemini': [
        ('gemini.google.com', 'Gemini'), ('aistudio.google.com', 'Google AI Studio'),
        ('notebooklm.google.com', 'NotebookLM'), ('bard.google.com', 'Bard (tên cũ của Gemini)'), ('labs.google', 'Google Labs (Flow, Whisk)'),
    ],
    'Microsoft Copilot': [
        ('copilot.microsoft.com', 'Microsoft Copilot'), ('copilot.cloud.microsoft', 'Microsoft 365 Copilot'),
        ('bing.com/chat', 'Chat AI trong Bing'), ('bing.com/images/create', 'Tạo ảnh AI của Bing'),
    ],
    'Claude, Perplexity, Grok': [
        ('claude.ai', 'Claude'), ('claude.com', 'Claude'), ('anthropic.com', 'Anthropic'),
        ('perplexity.ai', 'Perplexity'), ('grok.com', 'Grok'), ('x.ai', 'xAI'),
    ],
    'DeepSeek và AI Trung Quốc': [
        ('deepseek.com', 'DeepSeek'), ('kimi.com', 'Kimi'), ('kimi.ai', 'Kimi'), ('moonshot.cn', 'Moonshot'),
        ('qwen.ai', 'Qwen'), ('tongyi.aliyun.com', 'Tongyi Qianwen'), ('doubao.com', 'Doubao'),
        ('yiyan.baidu.com', 'Ernie Bot'), ('chatglm.cn', 'ChatGLM'), ('z.ai', 'Z.ai (GLM)'),
    ],
    'Mistral, Meta AI, Poe': [
        ('mistral.ai', 'Mistral, Le Chat'), ('meta.ai', 'Meta AI'), ('poe.com', 'Poe'),
        ('huggingface.co/chat', 'HuggingChat'), ('duck.ai', 'Duck.ai'), ('you.com', 'You.com'), ('phind.com', 'Phind'),
        ('lmarena.ai', 'LMArena'), ('openrouter.ai', 'OpenRouter'),
    ],
    'Nhập vai, bạn đồng hành AI': [
        ('character.ai', 'Character.AI'), ('janitorai.com', 'Janitor AI'), ('talkie-ai.com', 'Talkie'), ('replika.com', 'Replika'),
    ],
    'Viết bài và giải bài tập bằng AI': [
        ('quillbot.com', 'QuillBot viết lại câu'), ('gauthmath.com', 'Gauth giải toán'), ('brainly.com', 'Brainly'),
        ('jenni.ai', 'Jenni AI'), ('wordtune.com', 'Wordtune'), ('copy.ai', 'Copy.ai'), ('jasper.ai', 'Jasper'),
        ('writesonic.com', 'Writesonic'), ('undetectable.ai', 'Undetectable AI (qua mặt kiểm tra AI)'),
    ],
    'Tạo ảnh, video, âm thanh AI': [
        ('midjourney.com', 'Midjourney'), ('leonardo.ai', 'Leonardo'), ('ideogram.ai', 'Ideogram'),
        ('runwayml.com', 'Runway'), ('pika.art', 'Pika'), ('lumalabs.ai', 'Luma'), ('suno.com', 'Suno tạo nhạc'),
        ('elevenlabs.io', 'ElevenLabs giọng nói'), ('craiyon.com', 'Craiyon'), ('stability.ai', 'Stability AI'),
    ],
    'AI lập trình, giải code': [
        ('blackbox.ai', 'Blackbox AI'), ('codeium.com', 'Codeium / Windsurf'), ('tabnine.com', 'Tabnine'),
        ('cursor.com', 'Cursor'), ('v0.dev', 'v0 của Vercel'), ('bolt.new', 'Bolt'), ('lovable.dev', 'Lovable'),
        ('replit.com', 'Replit AI'), ('codium.ai', 'Qodo/CodiumAI'), ('deepai.org', 'DeepAI'),
    ],
    'AI làm web, slide, tài liệu (ít nổi)': [
        ('dora.run', 'Dora AI tạo web'), ('gamma.app', 'Gamma tạo slide'), ('tome.app', 'Tome'),
        ('beautiful.ai', 'Beautiful.ai'), ('chatpdf.com', 'ChatPDF'), ('humata.ai', 'Humata'),
        ('sidekickbrowser.com', 'Sidekick'), ('monica.im', 'Monica AI'), ('sider.ai', 'Sider AI'),
        ('merlin.foyer.work', 'Merlin'), ('theb.ai', 'TheB.AI'), ('flowgpt.com', 'FlowGPT'),
        ('perchance.org/ai-chat', 'Perchance AI chat'), ('venice.ai', 'Venice AI'),
    ],
}
AI_TOO_BROAD = {'google.com', 'microsoft.com', 'bing.com', 'live.com', 'office.com', 'windows.com', 'windowsupdate.com',
                'github.com', 'youtube.com', 'wikipedia.org', 'x.com', 'apple.com', 'amazon.com', 'facebook.com'}  # giống $aiTooBroad trong script

LISTS = {
    'keep': {
        'file': 'KeepAppList.xlsx', 'sheet': 'KeepAppList',
        'title': 'Ứng dụng được phép chạy (KeepAppList.xlsx)',
        'col_a': 'Tên chương trình hoặc đường dẫn thư mục',
        'apply': 'lockall-start',
        'hint': 'Dùng cho -lockall-start, -lockrun-start, -lockapp-start, -keepapp. Ghi tên tiến trình (chrome, mspaint) '
                'hoặc đường dẫn thư mục cài đặt (kết thúc bằng dấu \\). Mọi ứng dụng khác sẽ bị chặn.',
        'legend': [
            'HƯỚNG DẪN (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng 1 ứng dụng được phép chạy.',
            'Ghi tên tiến trình (ví dụ: chrome, mspaint) hoặc đường dẫn thư mục (có dấu \\ ở cuối).',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\KeepAppList.xlsx.',
            'Sửa danh sách xong, chạy lại -lockall-start để áp dụng.',
        ],
        'presets': {
            'Google Chrome': [
                ('chrome', 'Google Chrome (tên tiến trình)'),
                ('C:\\Program Files\\Google\\Chrome\\Application\\', 'Chrome - thư mục cài đặt (64-bit)'),
                ('C:\\Program Files (x86)\\Google\\Chrome\\Application\\', 'Chrome - thư mục cài đặt (32-bit)'),
                ('%LOCALAPPDATA%\\Google\\Chrome\\Application\\', 'Chrome - cài riêng cho từng người dùng'),
            ],
            'ONLYOFFICE': [
                ('DesktopEditors', 'ONLYOFFICE Desktop Editors (tên tiến trình)'),
                ('C:\\Program Files\\ONLYOFFICE\\DesktopEditors\\', 'ONLYOFFICE - thư mục cài đặt'),
            ],
            'Microsoft Paint': [('mspaint', 'Microsoft Paint')],
            'Word / Excel / PowerPoint': [
                ('winword', 'Microsoft Word'), ('excel', 'Microsoft Excel'), ('powerpnt', 'Microsoft PowerPoint'),
            ],
            'Notepad': [('notepad', 'Notepad')],
            'Microsoft Edge': [
                ('msedge', 'Microsoft Edge'),
                ('C:\\Program Files (x86)\\Microsoft\\Edge\\Application\\', 'Edge - thư mục cài đặt'),
            ],
        },
    },
    'allow': {
        'file': 'AllowPage.xlsx', 'sheet': 'AllowPage',
        'title': 'Trang web được phép (AllowPage.xlsx)',
        'col_a': 'Tên miền (ví dụ wikipedia.org)',
        'apply': 'allowpage',
        'hint': 'Dùng cho -allowpage: chặn internet, chỉ cho phép các trang trong danh sách. Có thể dán cả địa chỉ đầy đủ, '
                'script tự lấy tên miền. Lưu ý: hoạt động theo địa chỉ IP nên trang dùng nhiều tên miền/CDN có thể tải không đủ.',
        'legend': [
            'HƯỚNG DẪN - dùng cho lệnh -allowpage (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng 1 trang web được phép. Ghi tên miền, ví dụ: wikipedia.org',
            'Có thể dán cả địa chỉ đầy đủ (https://www.wikipedia.org/wiki/...): script tự lấy phần tên miền.',
            'Không cần ghi www. Script tự thử cả có www và không có www.',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\AllowPage.xlsx.',
            'Lưu ý: -allowpage hoạt động theo địa chỉ IP, trang dùng nhiều tên miền/CDN có thể tải không đủ.',
            'Sửa danh sách xong, chạy lại -allowpage để cập nhật.',
        ],
        'presets': {
            'Bộ trang học tập mẫu': [
                ('wikipedia.org', 'Wikipedia'),
                ('vietjack.vn', 'VietJack - bài giảng, giải bài tập'),
                ('olm.vn', 'OLM - học và luyện tập trực tuyến'),
                ('khanacademy.org', 'Khan Academy'),
            ],
        },
    },
    'block': {
        'file': 'BlockAppList.xlsx', 'sheet': 'BlockAppList',
        'title': 'Ứng dụng / thư mục bị chặn (BlockAppList.xlsx)',
        'col_a': 'Đường dẫn đầy đủ của file hoặc thư mục',
        'apply': 'blockapp',
        'hint': 'Dùng cho -blockapp: chặn không cho chạy các file/thư mục này. Phải là đường dẫn đầy đủ bắt đầu bằng ổ đĩa, '
                'hỗ trợ * và ?. Không dùng %LOCALAPPDATA% (lệnh chạy bằng SYSTEM), dùng C:\\Users\\*\\... thay thế.',
        'legend': [
            'HƯỚNG DẪN - dùng cho lệnh -blockapp (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng là đường dẫn ĐẦY ĐỦ của 1 file hoặc 1 thư mục cần chặn (phải bắt đầu bằng ổ đĩa, ví dụ C:\\).',
            'Hỗ trợ ký tự đại diện * và ?. Dòng không phải đường dẫn sẽ bị bỏ qua.',
            'Không dùng %LOCALAPPDATA% hay %USERPROFILE%: lệnh chạy bằng quyền SYSTEM nên chúng trỏ sai hồ sơ. Dùng C:\\Users\\*\\... thay thế.',
            'C:\\Windows, thư mục Veyon, C:\\cmd và ổ đĩa gốc luôn được bỏ qua để tránh làm hỏng máy.',
            'Không chặn được ứng dụng Microsoft Store (UWP). Gỡ chặn bằng -unblockapp.',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\BlockAppList.xlsx.',
        ],
        'presets': {
            'Steam': [('C:\\Program Files (x86)\\Steam', 'Steam - cả thư mục')],
            'Epic Games': [('C:\\Program Files\\Epic Games', 'Epic Games - cả thư mục')],
            'Roblox': [('C:\\Users\\*\\AppData\\Local\\Roblox', 'Roblox - cài trong hồ sơ từng học sinh')],
            'Thư mục game C:\\Games': [
                ('C:\\Games\\*.exe', 'Mọi file .exe trong C:\\Games'),
                ('C:\\Users\\Public\\Games', 'Thư mục game dùng chung'),
            ],
        },
    },
    'ai': {
        'file': 'BlockAIList.xlsx', 'sheet': 'BlockAIList',
        'title': 'Trang web và công cụ AI bị chặn (BlockAIList.xlsx)',
        'col_a': 'Tên miền AI (ví dụ chatgpt.com) hoặc tên miền/đường dẫn (bing.com/chat)',
        'apply': 'blockai',
        'hint': 'Dùng cho -blockai (nút "Khóa toàn bộ công cụ AI"). Mỗi dòng là một trang AI cần chặn. Tên miền chặn luôn cả tên miền con '
                '(ví dụ chatgpt.com chặn cả chat.chatgpt.com). Muốn chặn riêng một phần của trang lớn, ghi cả đường dẫn, ví dụ bing.com/chat. '
                'Ngoài danh sách này, nút còn tắt AI tích hợp sẵn trong Chrome, Edge, Firefox và Brave.',
        'legend': [
            'HƯỚNG DẪN - dùng cho lệnh -blockai (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng 1 trang AI cần chặn. Ghi tên miền, ví dụ: chatgpt.com',
            'Tên miền chặn luôn các tên miền con. Có thể ghi đường dẫn để chặn một phần của trang, ví dụ: bing.com/chat',
            'Không ghi tên miền quá rộng như google.com, microsoft.com (script tự bỏ qua để không chặn nhầm cả trang).',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\BlockAIList.xlsx.',
            'Sửa danh sách xong, bấm lại "Khóa toàn bộ công cụ AI" để cập nhật.',
            'Không chặn được AI nằm trong trang lớn (Google Docs, Canva, Quizlet...) và AI trong kết quả tìm kiếm Google.',
        ],
        'presets': AI_PRESETS,
    },
}


def read_list(path):
    """Đọc cột A (và ghi chú ở cột B) của sheet đầu tiên. Trả về danh sách (giá trị, ghi chú)."""
    from openpyxl import load_workbook  # nạp muộn: openpyxl khá nặng, chỉ cần khi đọc/ghi Excel
    rows = []
    if not os.path.exists(path):
        return rows
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None or str(a).strip() == '':
            continue
        b = ws.cell(r, 2).value
        rows.append((str(a).strip(), '' if b is None else str(b).strip()))
    return rows


def write_list(path, kind, rows):
    """Ghi file Excel đúng định dạng mà script đọc (cột A) kèm ghi chú ở cột B và hướng dẫn ở cột D."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, PatternFill, Side  # nạp muộn: openpyxl khá nặng, chỉ cần khi đọc/ghi Excel
    meta = LISTS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = meta['sheet']
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill('solid', fgColor='FFF2CC')
    for i, (a, b) in enumerate(rows, 1):
        ca = ws.cell(i, 1, a)
        ca.font = Font(name='Arial', size=11, color='0000FF')
        ca.fill = fill
        ca.border = border
        cb = ws.cell(i, 2, b)
        cb.font = Font(name='Arial', size=11)
        cb.border = border
    for i, t in enumerate(meta['legend'], 1):
        c = ws.cell(i, 4, t)
        c.font = Font(name='Arial', size=11, bold=(i == 1))
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 3
    ws.column_dimensions['D'].width = 90
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    tmp = path + '.tmp'
    wb.save(tmp)
    os.replace(tmp, path)


def check_entry(kind, value):
    """Trả về cảnh báo nếu script sẽ bỏ qua hoặc xử lý sai dòng này, ngược lại trả về chuỗi rỗng."""
    if kind == 'allow':
        d = re.sub(r'^[a-z]+://', '', value.lower())
        d = re.sub(r'[/?#:].*$', '', d)
        if not (re.match(r'^\d{1,3}(\.\d{1,3}){3}$', d) or re.match(r'^[a-z0-9-]+(\.[a-z0-9-]+)+$', d)):
            return 'không giống tên miền, script sẽ bỏ qua'
    elif kind == 'block':
        if not (re.match(r'^[A-Za-z]:\\', value) or value.startswith('\\\\')):
            return 'phải là đường dẫn đầy đủ bắt đầu bằng ổ đĩa, script sẽ bỏ qua'
        if re.search(r'%[^%]+%', value):
            return 'có biến môi trường, -blockapp chạy bằng SYSTEM nên biến sẽ trỏ sai hồ sơ'
    elif kind == 'ai':
        d = re.sub(r'^[a-z]+://', '', value.lower().strip())
        d = re.sub(r'[?#].*$', '', d).rstrip('/')
        d = re.sub(r'^www\.', '', d)
        if not re.match(r'^[a-z0-9-]+(\.[a-z0-9-]+)+(/[A-Za-z0-9._~%/-]*)?$', d):
            return 'không giống tên miền, script sẽ bỏ qua'
        if '/' not in d and d in AI_TOO_BROAD:
            return 'tên miền quá rộng, script sẽ bỏ qua để không chặn nhầm cả trang (thêm đường dẫn, ví dụ bing.com/chat)'
    return ''


# ---------------------------------------------------------------- bản riêng của danh sách (từ bản 3.8.0)
# Mỗi danh sách có thể có thêm các bản riêng, lưu cạnh bản chính với tên "<Tên>--<tên bản>.xlsx"
# (ví dụ AllowPage--KiemTraVan.xlsx). Chế độ lớp chọn bản riêng để đẩy xuống máy học sinh khi vào tiết.
VARIANT_SEP = '--'


def variant_clean(name):
    """Tên bản riêng an toàn để làm tên file: chữ, số, dấu cách, gạch ngang; tối đa 40 ký tự. Rỗng nếu không hợp lệ."""
    n = re.sub(r'[^\w \-]', '', str(name or ''), flags=re.UNICODE).replace(VARIANT_SEP, '-')
    return re.sub(r'\s+', ' ', n).strip(' -_')[:40]


def variant_file(kind, name):
    clean = variant_clean(name)
    if not clean:
        raise ValueError('Tên bản riêng không hợp lệ')
    stem = os.path.splitext(LISTS[kind]['file'])[0]
    return '%s%s%s.xlsx' % (stem, VARIANT_SEP, clean)


def list_variants(list_dir):
    """{loại danh sách: [tên bản riêng, ...]} theo các file có trong thư mục danh sách."""
    out = {k: [] for k in LISTS}
    try:
        names = os.listdir(list_dir)
    except OSError:
        return out
    for k, m in LISTS.items():
        prefix = os.path.splitext(m['file'])[0] + VARIANT_SEP
        for fn in sorted(names):
            if fn.startswith(prefix) and fn.lower().endswith('.xlsx') and not fn.startswith('~$'):
                v = fn[len(prefix):-5]
                if v and variant_clean(v) == v:
                    out[k].append(v)
    return out
