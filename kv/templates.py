# -*- coding: utf-8 -*-
"""File Excel mẫu (thời khóa biểu, danh sách lớp) và các bộ giờ tiết có sẵn."""

import io


def build_bells(am='07:00', pm='13:30', dur=45, brk=5, big_after=2, big_brk=20, per_half=5):
    """Sinh giờ 10 tiết: `per_half` tiết sáng bắt đầu lúc am, phần còn lại bắt đầu lúc pm; nghỉ brk phút, nghỉ lớn big_brk phút sau tiết big_after."""
    def add(hhmm, m):
        t = int(hhmm[:2]) * 60 + int(hhmm[3:5]) + m
        return '%02d:%02d' % (t // 60, t % 60)
    out = {}
    for half, start in ((0, am), (1, pm)):
        t = start
        for i in range(1, per_half + 1):
            n = half * per_half + i
            e = add(t, dur)
            out[n] = (t, e)
            t = add(e, big_brk if i == big_after else brk)
    return out


BELL_SETS = [
    ('mac-dinh', 'Sáng 7:00, chiều 13:30 · tiết 45 phút', dict(am='07:00', pm='13:30')),
    ('7h15', 'Sáng 7:15, chiều 13:30 · tiết 45 phút', dict(am='07:15', pm='13:30')),
    ('7h30', 'Sáng 7:30, chiều 13:00 · tiết 45 phút', dict(am='07:30', pm='13:00')),
    ('6h45', 'Sáng 6:45, chiều 13:15 · tiết 45 phút', dict(am='06:45', pm='13:15')),
    ('40p', 'Sáng 7:00, chiều 13:30 · tiết 40 phút', dict(am='07:00', pm='13:30', dur=40)),
    ('35p', 'Sáng 7:00, chiều 13:30 · tiết 35 phút', dict(am='07:00', pm='13:30', dur=35)),
]


def bell_set(key):
    for k, _label, kw in BELL_SETS:
        if k == key:
            return build_bells(**kw)
    return None


def bell_sets_info():
    return [{'key': k, 'label': lab, 'bells': {str(n): list(v) for n, v in build_bells(**kw).items()}} for k, lab, kw in BELL_SETS]


DAYS = ['THỨ HAI', 'THỨ BA', 'THỨ TƯ', 'THỨ NĂM', 'THỨ SÁU', 'THỨ BẢY']


def timetable_template(bells):
    """File .xlsx mẫu: sheet 'Thời khóa biểu' (đọc lại được bằng nút Nhập thời khóa biểu) và sheet 'Hướng dẫn'."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Thời khóa biểu'
    bold = Font(bold=True)
    ws['A1'], ws['B1'] = 'Giáo viên', ''
    ws['A2'], ws['B2'] = 'Môn', ''
    ws['A3'], ws['B3'] = 'Áp dụng từ', ''
    for r in (1, 2, 3):
        ws.cell(r, 1).font = bold
    ws['D3'] = '← có thể để trống; ngày dạng 15/09/2026'
    ws['D3'].font = Font(italic=True, color='888888')
    hdr = 5
    for j, t in enumerate(['Tiết', 'Bắt đầu', 'Kết thúc'] + DAYS, 1):
        c = ws.cell(hdr, j, t)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F5597')
        c.alignment = Alignment(horizontal='center')
    sample = {(1, 0): '10A1', (2, 0): '10A1', (3, 2): '11B2', (4, 2): '11B2', (1, 4): '12C3'}
    yellow = PatternFill('solid', fgColor='FFF2CC')
    for n in range(1, 11):
        r = hdr + n
        st, en = bells[n]
        ws.cell(r, 1, n).alignment = Alignment(horizontal='center')
        ws.cell(r, 2, st).alignment = Alignment(horizontal='center')
        ws.cell(r, 3, en).alignment = Alignment(horizontal='center')
        for d in range(len(DAYS)):
            c = ws.cell(r, 4 + d, sample.get((n, d), ''))
            c.alignment = Alignment(horizontal='center')
            if (n, d) in sample:
                c.fill = yellow
    ws.column_dimensions['A'].width = 12
    for col in 'BC':
        ws.column_dimensions[col].width = 11
    for col in 'DEFGHI':
        ws.column_dimensions[col].width = 12
    g = wb.create_sheet('Hướng dẫn')
    for i, line in enumerate([
        'CÁCH ĐIỀN',
        '1. Ô vàng là ví dụ, xóa hoặc thay bằng lớp của bạn. Ô trống nghĩa là không có tiết.',
        '2. Mỗi hàng là một tiết. Giờ ở hai cột "Bắt đầu" và "Kết thúc" quyết định lớp nào đang học; sửa cho đúng với trường.',
        '3. Có thể ghi "10A1" hoặc "Lớp 10A1 (P.201)": chương trình tự lấy tên lớp.',
        '4. Không đổi hàng tiêu đề (THỨ HAI … THỨ BẢY). Không cần dùng hết các cột thứ.',
        '5. Lưu file (.xlsx) rồi bấm "Nhập thời khóa biểu" trong trang Lớp & lịch dạy.',
        '6. Tên lớp nên trùng với tên lớp trong danh sách lớp (ví dụ 10A1), nếu khác chương trình sẽ báo ở trang Lớp & lịch dạy.',
    ], 1):
        g.cell(i, 1, line).font = Font(bold=(i == 1))
    g.column_dimensions['A'].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def classes_template():
    """File .xlsx mẫu danh sách lớp theo kiểu file điểm danh của trường (mỗi lớp một sheet, tên lớp ở E2, họ tên ở cột C)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    first = True
    for cls, names in (('10A1', ['Nguyễn Văn An', 'Trần Thị Bình', 'Lê Minh Châu']), ('10A2', ['Phạm Quốc Dũng', 'Hoàng Thị Em'])):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = cls
        ws['A1'] = 'DANH SÁCH LỚP (sửa tên lớp ở ô E2, họ tên từ cột C; mỗi lớp một sheet)'
        ws['D2'], ws['E2'] = 'Lớp:', cls
        ws['D2'].font = ws['E2'].font = Font(bold=True)
        ws['A4'], ws['B4'], ws['C4'] = 'STT', 'Mã HS', 'HỌ & TÊN'
        for c in ('A4', 'B4', 'C4'):
            ws[c].font = Font(bold=True)
        for i, n in enumerate(names, 1):
            ws.cell(4 + i, 1, i)
            ws.cell(4 + i, 3, n)
        ws.column_dimensions['C'].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
