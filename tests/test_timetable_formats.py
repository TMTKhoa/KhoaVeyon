# -*- coding: utf-8 -*-
import os, sys, tempfile, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from kv import school
from kv.runtime import ApiError


def build(header, rows, title='TKB'):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = title
    for j, v in enumerate(header): ws.cell(3, j + 1, v)
    for i, r in enumerate(rows):
        for j, v in enumerate(r): ws.cell(4 + i, j + 1, v)
    fd, p = tempfile.mkstemp(suffix='.xlsx'); os.close(fd); wb.save(p); return p


class TT(unittest.TestCase):
    def parse(self, header, rows, **k):
        p = build(header, rows, **k)
        try: return school.parse_timetable_xlsx(p)
        finally: os.remove(p)

    def test_two_time_columns_classic(self):
        d = self.parse(['', 'Bắt đầu', 'Kết thúc', 'THỨ HAI', 'THỨ BA'], [['', '07:00', '07:45', '10A1', '']])
        self.assertEqual(d['slots'], [{'day': 0, 'start': '07:00', 'end': '07:45', 'cls': '10A1'}])

    def test_short_day_names_and_range_in_one_cell(self):
        d = self.parse(['Giờ', 'T2', 'T3', 'T4'], [['7:00 - 7:45', '10A1', '', '10A2']])
        self.assertEqual([(s['day'], s['cls']) for s in d['slots']], [(0, '10A1'), (2, '10A2')])
        self.assertEqual(d['slots'][0]['end'], '07:45')

    def test_thu_so_with_date_and_class_with_room(self):
        d = self.parse(['Tiết', 'Thứ 2 (19/10)', 'Thứ 3 (20/10)'], [[1, 'Lớp 10A1 (P.201)', 'x'], [2, '', '11B2 - Tin']])
        got = {(s['day'], s['start']): s['cls'] for s in d['slots']}
        self.assertEqual(got[(0, '07:00')], '10A1'); self.assertEqual(got[(1, '07:50')], '11B2')
        self.assertTrue(d['default_bells'])

    def test_second_sheet_and_accent_free(self):
        wb = openpyxl.Workbook(); wb.active.title = 'Ghi chu'
        ws = wb.create_sheet('Lich'); ws.append(['', 'Bat dau', 'Ket thuc', 'Thu hai', 'Thu ba']); ws.append(['', '13:30', '14:15', '9A1', '9A2'])
        fd, p = tempfile.mkstemp(suffix='.xlsx'); os.close(fd); wb.save(p)
        try: d = school.parse_timetable_xlsx(p)
        finally: os.remove(p)
        self.assertEqual(len(d['slots']), 2)

    def test_clear_errors(self):
        with self.assertRaises(ApiError) as e: self.parse(['a', 'b'], [[1, 2]])
        self.assertIn('tiêu đề các thứ', str(e.exception))
        with self.assertRaises(ApiError) as e: self.parse(['', 'THỨ HAI', 'THỨ BA'], [['abc', '10A1', '']])
        self.assertIn('không có tiết nào', str(e.exception))


if __name__ == '__main__':
    unittest.main()
