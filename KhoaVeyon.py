# -*- coding: utf-8 -*-
"""
KhoaVeyon - công cụ trên máy giáo viên, giao diện là một trang web chạy ngay trên máy này:
  1. Sửa 3 danh sách Excel (KeepAppList / AllowPage / BlockAppList).
  2. Ra lệnh cho các máy học sinh qua Veyon (khóa máy, chặn web, chặn ứng dụng...).
  3. Phân phối file sang máy học sinh rồi chạy ứng dụng.

Chạy: python KhoaVeyon.py     (cần: pip install openpyxl)
"""
import json
import os
import re
import shutil
import subprocess
import sys
import threading

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side

VERSION = '1.2.0'
APP_TITLE = 'KhoaVeyon - Hỗ trợ quản lý lớp học'
APP_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, 'frozen', False) else __file__))
LEGACY_CONFIG_PATH = os.path.join(APP_DIR, 'KhoaVeyon.json')


def _appdata_config_path():
    base = os.environ.get('APPDATA')
    if os.name == 'nt' and base:
        return os.path.join(base, 'KhoaVeyon', 'KhoaVeyon.json')
    return ''


def _config_path():
    """Chế độ portable: cài đặt nằm trong thư mục data cạnh chương trình, chép cả thư mục sang máy khác là đủ.
    Nếu thư mục chương trình không ghi được (Program Files, USB khóa ghi) thì dùng %APPDATA%\\KhoaVeyon."""
    folder = os.path.join(APP_DIR, 'data')
    try:
        os.makedirs(folder, exist_ok=True)
        probe = os.path.join(folder, '.write_test')
        with open(probe, 'w') as f:
            f.write('x')
        os.remove(probe)
        return os.path.join(folder, 'KhoaVeyon.json')
    except Exception:  # noqa
        pass
    alt = _appdata_config_path()
    if alt:
        try:
            os.makedirs(os.path.dirname(alt), exist_ok=True)
            return alt
        except Exception:  # noqa
            pass
    return LEGACY_CONFIG_PATH


CONFIG_PATH = _config_path()


def abs_path(p):
    """Đường dẫn tương đối (ví dụ 'lists') được tính từ thư mục chứa chương trình, để chép cả thư mục là chạy được."""
    p = os.path.expandvars((p or '').strip())
    if not p:
        return p
    return p if os.path.isabs(p) or p.startswith('\\\\') else os.path.normpath(os.path.join(APP_DIR, p))


def find_veyon_cli(configured=''):
    """Tìm veyon-cli: đường dẫn đã cài đặt, thư mục 'veyon' cạnh chương trình, PATH, rồi Program Files."""
    cands = []
    if configured:
        cands.append(abs_path(configured))
    cands += [os.path.join(APP_DIR, 'veyon', 'veyon-cli.exe'), os.path.join(APP_DIR, 'veyon-cli.exe')]
    found = shutil.which('veyon-cli')
    if found:
        cands.append(found)
    for env in ('ProgramFiles', 'ProgramFiles(x86)', 'ProgramW6432'):
        base = os.environ.get(env)
        if base:
            cands.append(os.path.join(base, 'Veyon', 'veyon-cli.exe'))
    for c in cands:
        if c and os.path.isfile(c):
            return c
    return configured


def resource_path(name):
    """Đường dẫn tới file đi kèm (biểu tượng), cả khi chạy từ file .exe."""
    return os.path.join(getattr(sys, '_MEIPASS', APP_DIR), name)
CREATE_NO_WINDOW = 0x08000000 if os.name == 'nt' else 0

# (khóa, nhãn nút, tham số của script, tên mục mặc định trong Veyon "Programs & websites")
ACTIONS = [
    ('lockall-start', 'Khóa toàn bộ (chỉ app trong danh sách)', '-lockall-start', 'Khóa toàn bộ (chỉ app trong danh sách)'),
    ('lockall-stop', 'Mở khóa toàn bộ', '-lockall-stop', 'Mở khóa toàn bộ'),
    ('lockrun-start', 'Chỉ chặn mở app ngoài danh sách', '-lockrun-start', 'Chỉ chặn mở app ngoài danh sách'),
    ('lockrun-stop', 'Gỡ chặn mở app', '-lockrun-stop', 'Gỡ chặn mở app'),
    ('lockapp-start', 'Bật chế độ kiosk', '-lockapp-start', 'Bật chế độ kiosk'),
    ('lockapp-stop', 'Tắt chế độ kiosk', '-lockapp-stop', 'Tắt chế độ kiosk'),
    ('killapp', 'Đóng tất cả ứng dụng', '-killapp', 'Đóng tất cả ứng dụng'),
    ('keepapp', 'Đóng tất cả trừ danh sách', '-keepapp', 'Đóng tất cả trừ danh sách'),
    ('blockall', 'Chặn internet', '-blockall', 'Chặn internet'),
    ('allowpage', 'Chỉ cho phép trang trong danh sách', '-allowpage', 'Chỉ cho phép các trang trong danh sách'),
    ('allowall', 'Cho phép internet', '-allowall', 'Cho phép internet'),
    ('blockapp', 'Chặn ứng dụng trong danh sách', '-blockapp', 'Chặn ứng dụng'),
    ('unblockapp', 'Bỏ chặn ứng dụng', '-unblockapp', 'Bỏ chặn ứng dụng'),
    ('lockdown', 'Khóa Task Manager / Settings', '-lockdown', 'Khóa (Lockdown)'),
    ('unlock', 'Mở khóa Task Manager / Settings', '-unlock', 'Mở khóa (Unlock)'),
    ('mute', 'Tắt tiếng loa', '-mute', 'Tắt tiếng'),
    ('unmute', 'Bật tiếng loa', '-unmute', 'Bật tiếng'),
    ('clearhistory', 'Xóa lịch sử Chrome/Firefox', '-clearhistory', 'Xóa lịch sử trình duyệt'),
]
ACTION_BY_KEY = {a[0]: a for a in ACTIONS}
# Các lệnh có thể làm mất bài chưa lưu của học sinh, nên hỏi lại trước khi chạy
CONFIRM_KEYS = {'lockall-start', 'lockrun-start', 'lockapp-start', 'killapp', 'keepapp', 'clearhistory'}

DEFAULT_CONFIG = {
    'list_dir': 'lists',  # tương đối so với thư mục chương trình (portable); có thể đặt C:\\cmd
    'script_path': r'C:\cmd\KhoaVeyonServices.bat',
    'veyon_cli': r'C:\Program Files\Veyon\veyon-cli.exe',
    # Lệnh chạy một mục Veyon trên 1 máy. {cli} {host} {feature} sẽ được thay thế.
    'run_template': '"{cli}" feature start {host} "{feature}"',
    'list_template': '"{cli}" feature list {host}',
    'hosts': [],
    'sync_dir': '',
    'auto_sync': False,
    # Phân phối & chạy
    'dist_files': [],
    'dist_include_lists': True,
    'dist_target': r'C:\cmd',
    'dist_mode': 'schtasks',   # none | schtasks | veyon
    'dist_cmd': r'cmd.exe /c C:\cmd\KhoaVeyonServices.bat -lockall-start',
    'dist_action': 'lockall-start',
    'dist_interactive': False,
    'dist_payload_skip': [],   # file trong thư mục payload bị bỏ chọn khi phân phối
    'auto_push': True,   # bấm Lưu ở tab danh sách thì gửi file vừa sửa xuống máy học sinh
    'update_repo': 'TMTKhoa/KhoaVeyon',  # GitHub, dạng chủ-sở-hữu/tên-kho
    'update_check': True,  # tự kiểm tra bản mới mỗi lần khởi động (khi đã điền kho GitHub)
    'update_auto': False,  # tự cài và khởi động lại khi có bản mới (mặc định tắt, nên xem trước rồi bấm cập nhật)
    'auto_lock': False,  # sau khi thiết lập máy học sinh thì khóa C:\\cmd (học sinh chỉ đọc)
    'client_veyon_cli': r'C:\Program Files\Veyon\veyon-cli.exe',  # đường dẫn veyon-cli trên máy học sinh
    'vcfg_push': [],      # các khóa cấu hình Veyon đẩy xuống máy học sinh: [{'key','value'}]
    'vcfg_restart': True,  # khởi động lại Veyon Service sau khi đẩy cấu hình
    'netobj_location': 'KhoaVeyon',  # nhóm (location) trong danh sách máy của Veyon Master
    'shot_before_lock': False,  # chụp màn hình các máy trước khi chạy lệnh khóa / đóng ứng dụng
    'watch_repeat': '10',  # giây giữa hai lần chụp khi học sinh vẫn ở trang/ứng dụng vi phạm
    'watch_max': '5',      # số ảnh tối đa cho mỗi lần vi phạm
    'setup_done': False,  # đã chạy thiết lập ban đầu chưa (chưa thì mở trang Thiết lập khi khởi động)
    'features': {a[0]: a[3] for a in ACTIONS},
}

# ---------------------------------------------------------------- danh sách Excel
LISTS = {
    'keep': {
        'file': 'KeepAppList.xlsx', 'sheet': 'KeepAppList',
        'title': 'Ứng dụng được phép chạy (KeepAppList.xlsx)',
        'col_a': 'Tên chương trình hoặc đường dẫn thư mục',
        'apply': 'lockall-start',
        'hint': 'Dùng cho -lockall-start, -lockrun-start, -lockapp-start, -keepapp. Ghi tên tiến trình (chrome, mspaint) '
                'hoặc đường dẫn thư mục cài đặt (kết thúc bằng dấu \\). Mọi ứng dụng khác sẽ bị chặn.',
        'legend': [
            'HƯỚNG DẪN (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng 1 ứng dụng được phép chạy.',
            'Ghi tên tiến trình (ví dụ: chrome, mspaint) hoặc đường dẫn thư mục (có dấu \\ ở cuối).',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\KeepAppList.xlsx.',
            'Sửa danh sách xong, chạy lại -lockall-start để áp dụng.',
        ],
        'presets': {
            'Google Chrome': [
                ('chrome', 'Google Chrome (tên tiến trình)'),
                ('C:\\Program Files\\Google\\Chrome\\Application\\', 'Chrome - thư mục cài đặt (64-bit)'),
                ('C:\\Program Files (x86)\\Google\\Chrome\\Application\\', 'Chrome - thư mục cài đặt (32-bit)'),
                ('%LOCALAPPDATA%\\Google\\Chrome\\Application\\', 'Chrome - cài riêng cho từng người dùng'),
            ],
            'ONLYOFFICE': [
                ('DesktopEditors', 'ONLYOFFICE Desktop Editors (tên tiến trình)'),
                ('C:\\Program Files\\ONLYOFFICE\\DesktopEditors\\', 'ONLYOFFICE - thư mục cài đặt'),
            ],
            'Microsoft Paint': [('mspaint', 'Microsoft Paint')],
            'Word / Excel / PowerPoint': [
                ('winword', 'Microsoft Word'), ('excel', 'Microsoft Excel'), ('powerpnt', 'Microsoft PowerPoint'),
            ],
            'Notepad': [('notepad', 'Notepad')],
            'Microsoft Edge': [
                ('msedge', 'Microsoft Edge'),
                ('C:\\Program Files (x86)\\Microsoft\\Edge\\Application\\', 'Edge - thư mục cài đặt'),
            ],
        },
    },
    'allow': {
        'file': 'AllowPage.xlsx', 'sheet': 'AllowPage',
        'title': 'Trang web được phép (AllowPage.xlsx)',
        'col_a': 'Tên miền (ví dụ wikipedia.org)',
        'apply': 'allowpage',
        'hint': 'Dùng cho -allowpage: chặn internet, chỉ cho phép các trang trong danh sách. Có thể dán cả địa chỉ đầy đủ, '
                'script tự lấy tên miền. Lưu ý: hoạt động theo địa chỉ IP nên trang dùng nhiều tên miền/CDN có thể tải không đủ.',
        'legend': [
            'HƯỚNG DẪN - dùng cho lệnh -allowpage (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng 1 trang web được phép. Ghi tên miền, ví dụ: wikipedia.org',
            'Có thể dán cả địa chỉ đầy đủ (https://www.wikipedia.org/wiki/...): script tự lấy phần tên miền.',
            'Không cần ghi www. Script tự thử cả có www và không có www.',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\AllowPage.xlsx.',
            'Lưu ý: -allowpage hoạt động theo địa chỉ IP, trang dùng nhiều tên miền/CDN có thể tải không đủ.',
            'Sửa danh sách xong, chạy lại -allowpage để cập nhật.',
        ],
        'presets': {
            'Bộ trang học tập mẫu': [
                ('wikipedia.org', 'Wikipedia'),
                ('vietjack.vn', 'VietJack - bài giảng, giải bài tập'),
                ('olm.vn', 'OLM - học và luyện tập trực tuyến'),
                ('khanacademy.org', 'Khan Academy'),
            ],
        },
    },
    'block': {
        'file': 'BlockAppList.xlsx', 'sheet': 'BlockAppList',
        'title': 'Ứng dụng / thư mục bị chặn (BlockAppList.xlsx)',
        'col_a': 'Đường dẫn đầy đủ của file hoặc thư mục',
        'apply': 'blockapp',
        'hint': 'Dùng cho -blockapp: chặn không cho chạy các file/thư mục này. Phải là đường dẫn đầy đủ bắt đầu bằng ổ đĩa, '
                'hỗ trợ * và ?. Không dùng %LOCALAPPDATA% (lệnh chạy bằng SYSTEM), dùng C:\\Users\\*\\... thay thế.',
        'legend': [
            'HƯỚNG DẪN - dùng cho lệnh -blockapp (script chỉ đọc cột A, các cột khác chỉ để ghi chú)',
            'Cột A (ô màu vàng): mỗi dòng là đường dẫn ĐẦY ĐỦ của 1 file hoặc 1 thư mục cần chặn (phải bắt đầu bằng ổ đĩa, ví dụ C:\\).',
            'Hỗ trợ ký tự đại diện * và ?. Dòng không phải đường dẫn sẽ bị bỏ qua.',
            'Không dùng %LOCALAPPDATA% hay %USERPROFILE%: lệnh chạy bằng quyền SYSTEM nên chúng trỏ sai hồ sơ. Dùng C:\\Users\\*\\... thay thế.',
            'C:\\Windows, thư mục Veyon, C:\\cmd và ổ đĩa gốc luôn được bỏ qua để tránh làm hỏng máy.',
            'Không chặn được ứng dụng Microsoft Store (UWP). Gỡ chặn bằng -unblockapp.',
            'Không để dòng tiêu đề ở cột A. Lưu file thành C:\\cmd\\BlockAppList.xlsx.',
        ],
        'presets': {
            'Steam': [('C:\\Program Files (x86)\\Steam', 'Steam - cả thư mục')],
            'Epic Games': [('C:\\Program Files\\Epic Games', 'Epic Games - cả thư mục')],
            'Roblox': [('C:\\Users\\*\\AppData\\Local\\Roblox', 'Roblox - cài trong hồ sơ từng học sinh')],
            'Thư mục game C:\\Games': [
                ('C:\\Games\\*.exe', 'Mọi file .exe trong C:\\Games'),
                ('C:\\Users\\Public\\Games', 'Thư mục game dùng chung'),
            ],
        },
    },
}


def read_list(path):
    """Đọc cột A (và ghi chú ở cột B) của sheet đầu tiên. Trả về danh sách (giá trị, ghi chú)."""
    rows = []
    if not os.path.exists(path):
        return rows
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None or str(a).strip() == '':
            continue
        b = ws.cell(r, 2).value
        rows.append((str(a).strip(), '' if b is None else str(b).strip()))
    return rows


def write_list(path, kind, rows):
    """Ghi file Excel đúng định dạng mà script đọc (cột A) kèm ghi chú ở cột B và hướng dẫn ở cột D."""
    meta = LISTS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = meta['sheet']
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill('solid', fgColor='FFF2CC')
    for i, (a, b) in enumerate(rows, 1):
        ca = ws.cell(i, 1, a)
        ca.font = Font(name='Arial', size=11, color='0000FF')
        ca.fill = fill
        ca.border = border
        cb = ws.cell(i, 2, b)
        cb.font = Font(name='Arial', size=11)
        cb.border = border
    for i, t in enumerate(meta['legend'], 1):
        c = ws.cell(i, 4, t)
        c.font = Font(name='Arial', size=11, bold=(i == 1))
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 3
    ws.column_dimensions['D'].width = 90
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    tmp = path + '.tmp'
    wb.save(tmp)
    os.replace(tmp, path)


def check_entry(kind, value):
    """Trả về cảnh báo nếu script sẽ bỏ qua hoặc xử lý sai dòng này, ngược lại trả về chuỗi rỗng."""
    if kind == 'allow':
        d = re.sub(r'^[a-z]+://', '', value.lower())
        d = re.sub(r'[/?#:].*$', '', d)
        if not (re.match(r'^\d{1,3}(\.\d{1,3}){3}$', d) or re.match(r'^[a-z0-9-]+(\.[a-z0-9-]+)+$', d)):
            return 'không giống tên miền, script sẽ bỏ qua'
    elif kind == 'block':
        if not (re.match(r'^[A-Za-z]:\\', value) or value.startswith('\\\\')):
            return 'phải là đường dẫn đầy đủ bắt đầu bằng ổ đĩa, script sẽ bỏ qua'
        if re.search(r'%[^%]+%', value):
            return 'có biến môi trường, -blockapp chạy bằng SYSTEM nên biến sẽ trỏ sai hồ sơ'
    return ''


# ---------------------------------------------------------------- cấu hình và lệnh
def load_config():
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    path = CONFIG_PATH
    old_json = os.path.join(os.path.dirname(CONFIG_PATH), 'VeyonHelper.json')  # tên cũ của chương trình
    if not os.path.exists(path) and os.path.exists(old_json):
        path = old_json
    if not os.path.exists(path):
        old = _appdata_config_path()
        path = old if old and os.path.exists(old) else LEGACY_CONFIG_PATH
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        for k, v in data.items():
            if k == 'features' and isinstance(v, dict):
                cfg['features'].update(v)
            else:
                cfg[k] = v
    except FileNotFoundError:
        pass
    except Exception:
        pass
    return cfg


def save_config(cfg):
    with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def client_command(cfg, switch):
    """Lệnh để dán vào Veyon (Run program)."""
    return 'cmd.exe /c %s %s' % (cfg['script_path'], switch)


def veyon_command(cfg, host, feature, template_key='run_template'):
    return cfg[template_key].format(cli=find_veyon_cli(cfg['veyon_cli']), host=host, feature=feature)


def run_shell(cmdline, timeout=60):
    try:
        p = subprocess.run(cmdline, shell=True, capture_output=True, text=True, errors='replace',
                           timeout=timeout, creationflags=CREATE_NO_WINDOW)
        out = ((p.stdout or '') + (p.stderr or '')).strip()
        return p.returncode, out
    except subprocess.TimeoutExpired:
        return -1, 'Quá thời gian chờ'
    except Exception as e:  # noqa
        return -1, str(e)


def run_args(args, timeout=60):
    """Chạy lệnh dạng danh sách tham số, không cho đọc bàn phím (tránh treo khi lệnh hỏi mật khẩu)."""
    try:
        p = subprocess.run(args, capture_output=True, text=True, errors='replace', timeout=timeout,
                           stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW)
        return p.returncode, ((p.stdout or '') + (p.stderr or '')).strip()
    except subprocess.TimeoutExpired:
        return -1, 'Quá thời gian chờ'
    except Exception as e:  # noqa
        return -1, str(e)


def unc_target(host, target_dir):
    """C:\\cmd trên máy học sinh -> \\\\HOST\\C$\\cmd (chia sẻ quản trị)."""
    m = re.match(r'^([A-Za-z]):[\\/]*(.*)$', (target_dir or '').strip())
    if not m:
        raise ValueError('Thư mục đích phải bắt đầu bằng ổ đĩa, ví dụ C:\\cmd')
    return '\\\\%s\\%s$\\%s' % (host, m.group(1).upper(), m.group(2).replace('/', '\\'))


def distribute_to_host(host, files, target_dir):
    """Chép file/thư mục sang máy học sinh. Trả về số mục đã chép."""
    dst = unc_target(host, target_dir)
    os.makedirs(dst, exist_ok=True)
    n = 0
    for f in files:
        src = abs_path(f)
        if os.path.isdir(src):
            shutil.copytree(src, os.path.join(dst, os.path.basename(src.rstrip('\\/'))), dirs_exist_ok=True)
        else:
            shutil.copy2(src, os.path.join(dst, os.path.basename(src)))
        n += 1
    return n


def logged_on_user(host):
    rc, out = run_args(['query', 'user', '/server:' + host], timeout=20)
    if rc != 0:
        return ''
    for line in out.splitlines()[1:]:
        t = line.replace('>', ' ').split()
        if t:
            return t[0]
    return ''


def remote_run(host, cmdline, interactive=False, name='KhoaVeyonRun'):
    """Chạy 1 lệnh trên máy học sinh bằng Task Scheduler (schtasks) rồi kích hoạt ngay.
    Mặc định chạy bằng SYSTEM (nền, học sinh không thấy cửa sổ). interactive=True: chạy bằng người đang đăng nhập."""
    ru = 'SYSTEM'
    if interactive:
        ru = logged_on_user(host)
        if not ru:
            return -1, 'Không tìm thấy người đang đăng nhập trên máy này'
    create = ['schtasks', '/Create', '/S', host, '/TN', name, '/TR', cmdline, '/SC', 'ONCE', '/ST', '00:00',
              '/RU', ru, '/F']
    create += ['/IT'] if interactive else ['/RL', 'HIGHEST']
    rc, out = run_args(create)
    if rc != 0:
        return rc, out
    rc2, out2 = run_args(['schtasks', '/Run', '/S', host, '/TN', name])
    return rc2, (out + '\n' + out2).strip()


def sync_lists(cfg):
    """Chép 3 file Excel từ thư mục danh sách sang thư mục đồng bộ. Trả về (số file đã chép, lỗi)."""
    dst = (cfg.get('sync_dir') or '').strip()
    if not dst:
        return 0, 'Chưa đặt thư mục đồng bộ'
    dst = abs_path(dst)
    os.makedirs(dst, exist_ok=True)
    n = 0
    for meta in LISTS.values():
        src = os.path.join(abs_path(cfg['list_dir']), meta['file'])
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(dst, meta['file']))
            n += 1
    return n, ''

# ---------------------------------------------------------------- máy chủ web cục bộ
# Giao diện là một trang web chạy tại http://127.0.0.1:<cổng>/ (chỉ máy này truy cập được), có mã bảo vệ
# ngẫu nhiên mỗi lần chạy nên trang web khác không thể ra lệnh qua trình duyệt của giáo viên.
import http.server
import secrets
import time
import traceback
import urllib.parse
import urllib.error
import urllib.request
import webbrowser
from concurrent.futures import ThreadPoolExecutor, as_completed

WEB_DIR = os.path.join(APP_DIR, 'web')
PAYLOAD_DIR = os.path.join(APP_DIR, 'payload')
IDLE_EXIT_SECONDS = 180      # đóng cửa sổ trình duyệt thì chương trình tự thoát sau chừng này giây
MAX_UPLOAD = 500 * 1024 * 1024

STATE = {'token': '', 'port': 0, 'last_ping': time.time(), 'pinged': False, 'busy': 0, 'server': None}
CFG_LOCK = threading.RLock()
CFG = load_config()
LOG = []
LOG_LOCK = threading.Lock()
LOG_SEQ = [0]

GROUPS = [
    ('Khóa máy', ['lockall-start', 'lockall-stop', 'lockrun-start', 'lockrun-stop', 'lockapp-start', 'lockapp-stop']),
    ('Ứng dụng', ['killapp', 'keepapp', 'blockapp', 'unblockapp']),
    ('Internet', ['blockall', 'allowpage', 'allowall']),
    ('Hệ thống và âm thanh', ['lockdown', 'unlock', 'mute', 'unmute', 'clearhistory']),
]
ICONS = {
    'lockall-start': '🔒', 'lockall-stop': '🔓', 'lockrun-start': '⛔', 'lockrun-stop': '✅',
    'lockapp-start': '🖥️', 'lockapp-stop': '🔙', 'killapp': '✖️', 'keepapp': '🧹',
    'blockall': '🚫', 'allowpage': '🌐', 'allowall': '📶', 'blockapp': '🎮', 'unblockapp': '✅',
    'lockdown': '🛡️', 'unlock': '🗝️', 'mute': '🔇', 'unmute': '🔊', 'clearhistory': '🗑️',
}
# Nút cam = bật hạn chế, nút xanh lá = gỡ hạn chế, nút xám = việc khác
LOCK_KEYS = {'lockall-start', 'lockrun-start', 'lockapp-start', 'blockall', 'allowpage', 'blockapp',
             'lockdown', 'killapp', 'keepapp', 'clearhistory'}
FREE_KEYS = {'lockall-stop', 'lockrun-stop', 'lockapp-stop', 'allowall', 'unblockapp', 'unlock'}

STR_KEYS = ['list_dir', 'script_path', 'veyon_cli', 'run_template', 'list_template', 'sync_dir',
            'dist_target', 'dist_mode', 'dist_cmd', 'dist_action', 'update_repo', 'watch_repeat', 'watch_max']
BOOL_KEYS = ['auto_sync', 'auto_push', 'dist_include_lists', 'dist_interactive', 'setup_done', 'auto_lock', 'update_check', 'update_auto', 'shot_before_lock']
LIST_KEYS = ['dist_files', 'dist_payload_skip']


class ApiError(Exception):
    pass


def log(msg, kind=''):
    with LOG_LOCK:
        LOG_SEQ[0] += 1
        LOG.append({'id': LOG_SEQ[0], 't': time.strftime('%H:%M:%S'), 'msg': msg, 'kind': kind})
        if len(LOG) > 1500:
            del LOG[:300]


def start_job(fn, *args):
    def run():
        STATE['busy'] += 1
        try:
            fn(*args)
        except Exception as e:  # noqa
            log('LỖI không mong muốn: %s' % e, 'err')
        finally:
            STATE['busy'] -= 1
    threading.Thread(target=run, daemon=True).start()


def pool_run(fn, hosts):
    ok_n = err_n = 0
    with ThreadPoolExecutor(max_workers=min(8, max(1, len(hosts)))) as ex:
        futs = {ex.submit(fn, h): h for h in hosts}
        for fut in as_completed(futs):
            h = futs[fut]
            try:
                ok, text = fut.result()
            except Exception as e:  # noqa
                ok, text = False, str(e)
            if ok:
                ok_n += 1
            else:
                err_n += 1
            log('[%s] %s %s' % (h, 'OK' if ok else 'LỖI', text), 'ok' if ok else 'err')
    log('--- xong: %d thành công, %d lỗi ---' % (ok_n, err_n), 'warn' if err_n else 'head')


# ---------------------------------------------------------------- các công việc chạy nền
def veyon_feature(key):
    with CFG_LOCK:
        return CFG['features'].get(key) or ACTION_BY_KEY[key][3]


def job_action(key, hosts):
    label = ACTION_BY_KEY[key][1]
    feature = veyon_feature(key)
    log('>>> %s  (mục Veyon: "%s", %d máy)' % (label, feature, len(hosts)), 'head')

    def one(h):
        try:
            with CFG_LOCK:
                cmd = veyon_command(CFG, h, feature)
        except Exception as e:  # noqa
            return False, 'Lỗi mẫu lệnh: %s' % e
        rc, out = run_shell(cmd)
        return rc == 0, out if rc == 0 else '(%s) %s' % (rc, out)
    pool_run(one, hosts)


def job_dist(hosts, files, mode, cmd, feature, interactive, target, title):
    log('>>> %s  (%d máy)' % (title, len(hosts)), 'head')

    def one(h):
        msgs = []
        if files:
            try:
                n = distribute_to_host(h, files, target)
                msgs.append('đã chép %d mục' % n)
            except Exception as e:  # noqa
                return False, 'lỗi chép file: %s%s' % (e, ' (bỏ qua bước chạy)' if mode != 'none' else '')
        if mode == 'schtasks':
            rc, out = remote_run(h, cmd, interactive)
            if rc != 0:
                return False, '%schạy lệnh lỗi (%s) %s' % (', '.join(msgs) + ', ' if msgs else '', rc, out)
            msgs.append('đã chạy lệnh')
        elif mode == 'veyon':
            try:
                with CFG_LOCK:
                    c = veyon_command(CFG, h, feature)
            except Exception as e:  # noqa
                return False, 'Lỗi mẫu lệnh: %s' % e
            rc, out = run_shell(c)
            if rc != 0:
                return False, '%schạy mục Veyon lỗi (%s) %s' % (', '.join(msgs) + ', ' if msgs else '', rc, out)
            msgs.append('đã kích hoạt mục Veyon')
        return True, ', '.join(msgs)
    pool_run(one, hosts)


def job_after_save(hosts, path, push, apply_key, target):
    """Sau khi lưu danh sách: gửi file xuống máy học sinh (nếu bật), xong mới áp dụng lệnh để máy đọc đúng bản mới."""
    if push:
        job_dist(hosts, [path], 'none', '', '', False, target, 'Gửi %s xuống máy học sinh' % os.path.basename(path))
    if apply_key:
        job_action(apply_key, hosts)



# ---------------------------------------------------------------- thiết lập ban đầu (tự động)
import socket
import uuid

VEYON_PORT = 11100
STUDENT_SETUP_BAT = 'Setup-MayHocSinh.bat'
SETUP_RESULT = lambda: os.path.join(os.path.dirname(CONFIG_PATH), 'setup-result.json')  # noqa
DEFAULT_PRESETS = {  # danh sách mẫu tạo sẵn ở lần đầu (chỉ khi file chưa tồn tại)
    'keep': ['Google Chrome', 'ONLYOFFICE', 'Microsoft Paint'],
    'allow': ['Bộ trang học tập mẫu'],
    'block': ['Steam', 'Epic Games', 'Roblox', 'Thư mục game C:\\Games'],
}

STUDENT_BAT = r"""@echo off
chcp 65001 >nul
REM KhoaVeyon - thiet lap may hoc sinh. Chay 1 lan, chuot phai > Run as administrator.
net session >nul 2>&1
if errorlevel 1 (
  echo Hay chay file nay bang quyen quan tri: chuot phai - Run as administrator.
  pause
  exit /b 1
)
if not exist "C:\cmd" mkdir "C:\cmd"
REM Cho phep chia se file va quan ly Task Scheduler tu xa (dung ca ten nhom tieng Anh lan ma ngon ngu-doc-lap)
netsh advfirewall firewall set rule group="File and Printer Sharing" new enable=Yes >nul 2>&1
netsh advfirewall firewall set rule group="@FirewallAPI.dll,-28502" new enable=Yes >nul 2>&1
netsh advfirewall firewall set rule group="Remote Scheduled Tasks Management" new enable=Yes >nul 2>&1
netsh advfirewall firewall set rule group="@FirewallAPI.dll,-33502" new enable=Yes >nul 2>&1
REM Cho phep tai khoan quan tri cuc bo dung chia se C$ qua mang (may khong vao domain)
reg add "HKLM\SOFTWARE\Microsoft\Windows\CurrentVersion\Policies\System" /v LocalAccountTokenFilterPolicy /t REG_DWORD /d 1 /f >nul
netsh advfirewall firewall add rule name="Veyon Server" dir=in action=allow protocol=TCP localport=11100 profile=any >nul 2>&1
REM Bao dam dich vu Veyon tu chay
sc config VeyonService start= auto >nul 2>&1
sc start VeyonService >nul 2>&1
echo.
echo Xong. May hoc sinh san sang nhan file va lenh tu KhoaVeyon.
pause
"""


def is_admin():
    if os.name != 'nt':
        return True
    try:
        import ctypes
        return bool(ctypes.windll.shell32.IsUserAnAdmin())  # noqa
    except Exception:  # noqa
        return False


def find_script_file(cfg):
    """Tìm KhoaVeyonServices.bat để tự chép vào payload (thư mục payload, cạnh chương trình, thư mục cha, C:\\cmd)."""
    name = os.path.basename((cfg.get('script_path') or '').replace('\\', '/')) or 'KhoaVeyonServices.bat'
    cands = [os.path.join(PAYLOAD_DIR, name), os.path.join(APP_DIR, name), os.path.join(os.path.dirname(APP_DIR), name),
             os.path.join(os.path.dirname(os.path.dirname(APP_DIR)), name), cfg.get('script_path') or '',
             os.path.join('C:\\cmd', name)]
    for c in cands:
        if c and os.path.isfile(c):
            return c
    return ''


def seed_default_lists(cfg):
    """Tạo 3 file Excel mẫu nếu chưa có. Trả về số file đã tạo."""
    made = 0
    for kind, names in DEFAULT_PRESETS.items():
        path = os.path.join(abs_path(cfg['list_dir']), LISTS[kind]['file'])
        if os.path.exists(path):
            continue
        rows = []
        for n in names:
            rows += LISTS[kind]['presets'].get(n, [])
        write_list(path, kind, rows)
        made += 1
    return made


def ensure_payload(cfg):
    """Tạo Setup-MayHocSinh.bat trong payload và chép KhoaVeyonServices.bat vào payload nếu tìm thấy. Trả về (đã tạo, thiếu)."""
    os.makedirs(PAYLOAD_DIR, exist_ok=True)
    with open(os.path.join(PAYLOAD_DIR, STUDENT_SETUP_BAT), 'w', encoding='utf-8', newline='') as f:
        f.write(STUDENT_BAT.replace('\r\n', '\n').replace('\n', '\r\n'))
    name = os.path.basename((cfg.get('script_path') or '').replace('\\', '/')) or 'KhoaVeyonServices.bat'
    dst = os.path.join(PAYLOAD_DIR, name)
    got = []
    if not os.path.isfile(dst):
        src = find_script_file(cfg)
        if src:
            shutil.copy2(src, dst)
            got.append(name)
    return got, ([] if os.path.isfile(dst) else [name])


def desired_programs(cfg):
    """Các mục 'Programs & websites' cần có trong Veyon: (tên mục, dòng lệnh)."""
    return [(cfg['features'].get(a[0]) or a[3], 'cmd.exe /c %s %s' % (cfg['script_path'], a[2])) for a in ACTIONS]


def _veyon_get_programs(cli):
    rc, out = run_args([cli, 'config', 'get', 'DesktopServices/PredefinedPrograms'], timeout=30)
    if rc != 0:
        return None, out
    out = out.strip()
    if not out:
        return [], ''
    try:
        data = json.loads(out)
        if isinstance(data, str):
            data = json.loads(data)
        return (data if isinstance(data, list) else []), ''
    except Exception as e:  # noqa
        return None, 'Không đọc được kết quả veyon-cli: %s' % e


def veyon_items_status(cfg):
    """Trả về (số mục đã có, tổng, lỗi). Không ghi gì."""
    cli = find_veyon_cli(cfg['veyon_cli'])
    if not os.path.isfile(cli):
        return 0, len(ACTIONS), 'Chưa thấy veyon-cli'
    cur, err = _veyon_get_programs(cli)
    if cur is None:
        return 0, len(ACTIONS), err
    have = {str(x.get('Name', '')): str(x.get('Path', '')) for x in cur if isinstance(x, dict)}
    ok = sum(1 for n, c in desired_programs(cfg) if have.get(n) == c)
    return ok, len(ACTIONS), ''


def apply_veyon_items(cfg):
    """Ghi các mục vào cấu hình Veyon (giữ nguyên mục khác, cập nhật mục trùng tên). Cần quyền quản trị. Trả về (ok, thông báo)."""
    cli = find_veyon_cli(cfg['veyon_cli'])
    if not os.path.isfile(cli):
        return False, 'Chưa thấy veyon-cli.exe (đặt đường dẫn ở Cài đặt)'
    cur, err = _veyon_get_programs(cli)
    if cur is None:
        return False, 'Không đọc được cấu hình Veyon hiện tại: %s' % err
    try:  # sao lưu trước khi sửa
        bk = os.path.join(os.path.dirname(CONFIG_PATH), 'veyon-programs-backup.json')
        with open(bk, 'w', encoding='utf-8') as f:
            json.dump(cur, f, ensure_ascii=False, indent=2)
    except OSError:
        pass
    items = [x for x in cur if isinstance(x, dict)]
    by_name = {str(x.get('Name', '')): x for x in items}
    added = updated = 0
    for name, path in desired_programs(cfg):
        x = by_name.get(name)
        if x is None:
            items.append({'Name': name, 'Path': path, 'Uid': '{%s}' % uuid.uuid4()})
            added += 1
        elif x.get('Path') != path:
            x['Path'] = path
            updated += 1
    rc, out = run_args([cli, 'config', 'set', 'DesktopServices/PredefinedPrograms',
                        json.dumps(items, ensure_ascii=False, separators=(',', ':'))], timeout=60)
    if rc != 0:
        return False, 'veyon-cli config set lỗi (%s): %s' % (rc, out)
    again, err = _veyon_get_programs(cli)
    have = {str(x.get('Name', '')) for x in (again or []) if isinstance(x, dict)}
    miss = [n for n, _ in desired_programs(cfg) if n not in have]
    if miss:
        return False, 'Đã ghi nhưng Veyon chưa nhận đủ mục (thiếu %d). Có thể bản Veyon này lưu khác, hãy tạo thủ công theo README.' % len(miss)
    return True, 'Đã thêm %d mục, cập nhật %d mục. Hãy mở lại Veyon Master để thấy các mục mới.' % (added, updated)


def elevated_apply_veyon_items(cfg):
    """Gọi apply_veyon_items với quyền quản trị: nếu chưa là admin thì bật hộp thoại UAC chạy lại chương trình ở chế độ riêng."""
    if is_admin():
        return apply_veyon_items(cfg)
    try:
        os.remove(SETUP_RESULT())
    except OSError:
        pass
    import ctypes
    exe = sys.executable
    params = '"%s" --apply-veyon-items' % os.path.abspath(__file__)
    r = ctypes.windll.shell32.ShellExecuteW(None, 'runas', exe, params, APP_DIR, 0)  # noqa
    if r <= 32:
        return False, 'Bạn đã từ chối quyền quản trị (hoặc không mở được UAC)'
    for _ in range(120):
        time.sleep(0.5)
        if os.path.exists(SETUP_RESULT()):
            time.sleep(0.2)
            try:
                with open(SETUP_RESULT(), 'r', encoding='utf-8') as f:
                    d = json.load(f)
                return bool(d.get('ok')), d.get('msg', '')
            except Exception:  # noqa
                continue
    return False, 'Quá thời gian chờ kết quả từ cửa sổ quản trị'


def local_ips():
    ips = set()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('10.255.255.255', 1))
        ips.add(s.getsockname()[0])
        s.close()
    except Exception:  # noqa
        pass
    try:
        ips.update(socket.gethostbyname_ex(socket.gethostname())[2])
    except Exception:  # noqa
        pass
    return {i for i in ips if i and not i.startswith('127.')}


def tcp_open(host, port, timeout=0.5):
    try:
        with socket.create_connection((host, port), timeout=timeout):
            return True
    except OSError:
        return False


def scan_veyon_hosts(prefix=None, timeout=0.5):
    """Quét dải /24 của mạng hiện tại tìm máy đang mở cổng Veyon (11100). Trả về danh sách tên (hoặc IP)."""
    mine = local_ips()
    if prefix is None:
        if not mine:
            return [], 'Không xác định được địa chỉ mạng của máy này'
        prefix = sorted(mine)[0].rsplit('.', 1)[0]
    found = []

    def probe(i):
        ip = '%s.%d' % (prefix, i)
        return ip if ip not in mine and tcp_open(ip, VEYON_PORT, timeout) else None
    with ThreadPoolExecutor(max_workers=96) as ex:
        for ip in ex.map(probe, range(1, 255)):
            if ip:
                found.append(ip)
    names = []
    for ip in found:
        name = ip
        try:
            n = socket.gethostbyaddr(ip)[0].split('.')[0]
            if n and socket.gethostbyname(n) == ip:
                name = n.upper()
        except Exception:  # noqa
            pass
        names.append(name)
    return names, ''


def diagnose_host(cfg, host):
    """Kiểm tra từng điều kiện để KhoaVeyon dùng được với máy này. Trả về (danh sách (nhãn, đạt, gợi ý), ghi chú)."""
    checks = []
    veyon_ok = tcp_open(host, VEYON_PORT, 2)
    checks.append(('Veyon Service (cổng %d)' % VEYON_PORT, veyon_ok, 'Cài/bật Veyon Service trên máy học sinh, kiểm tra tường lửa'))
    try:
        unc = unc_target(host, cfg['dist_target'])
        share_ok = os.path.isdir(unc) or os.path.isdir(os.path.dirname(unc.rstrip('\\')))
    except Exception:  # noqa
        unc, share_ok = '', False
    checks.append(('Chia sẻ C$ (chép file)', share_ok,
                   'Chạy %s trên máy học sinh (quyền quản trị) và đăng nhập máy này bằng tài khoản quản trị của máy đó' % STUDENT_SETUP_BAT))
    rc, _ = run_args(['schtasks', '/Query', '/S', host, '/FO', 'LIST'], timeout=25)
    checks.append(('Task Scheduler từ xa (chạy lệnh)', rc == 0, 'Cần bật nhóm tường lửa "Remote Scheduled Tasks Management" (có trong %s)' % STUDENT_SETUP_BAT))
    sname = os.path.basename((cfg.get('script_path') or '').replace('\\', '/'))
    has = bool(unc) and os.path.isfile(os.path.join(unc, sname))
    checks.append(('Đã có %s trên máy' % sname, has, 'Bấm "Kiểm tra và gửi file xuống máy" để chép'))
    rc, _ = run_args(['sc', '\\\\' + host, 'query', 'KhoaVeyonService_blockall'], timeout=20)
    checks.append(('Đã cài dịch vụ (-install)', rc == 0, 'Sẽ tự cài sau khi chép file (cần chia sẻ và Task Scheduler từ xa)'))
    return checks


CMDLOCK_NAME = 'CmdLock.bat'
CMDLOCK_BAT = r"""@echo off
REM KhoaVeyon: khoa / mo khoa thu muc chua file nay (C:\cmd). Chay bang SYSTEM hoac quan tri.
REM Dung SID nen dung cho moi ngon ngu Windows: S-1-5-18 SYSTEM, S-1-5-32-544 Administrators, S-1-5-32-545 Users
set "D=%~dp0"
if "%D:~-1%"=="\" set "D=%D:~0,-1%"
if /i "%~1"=="unlock" goto UNLOCK
if not exist "%D%\log" mkdir "%D%\log"
icacls "%D%" /inheritance:r /grant:r "*S-1-5-18:(OI)(CI)F" "*S-1-5-32-544:(OI)(CI)F" "*S-1-5-32-545:(OI)(CI)RX" /T /C >nul
if errorlevel 1 exit /b 1
REM Thu muc log: hoc sinh van ghi duoc de nhat ky hoat dong chay tiep
icacls "%D%\log" /inheritance:r /grant:r "*S-1-5-18:(OI)(CI)F" "*S-1-5-32-544:(OI)(CI)F" "*S-1-5-32-545:(OI)(CI)M" /T /C >nul
exit /b %errorlevel%
:UNLOCK
icacls "%D%" /reset /T /C >nul
exit /b %errorlevel%
"""


def parse_task_result(csv_text):
    """Lấy mã 'Last Result' (cột 7) từ kết quả schtasks /Query /FO CSV /V. Trả về số nguyên hoặc None."""
    import csv
    import io
    rows = list(csv.reader(io.StringIO(csv_text)))
    for r in rows[1:] if len(rows) > 1 else []:
        if len(r) > 6:
            try:
                return int(r[6])
            except ValueError:
                return None
    return None


def remote_run_wait(host, cmdline, timeout=60, interactive=False):
    """Như remote_run (mặc định SYSTEM) nhưng đợi lệnh chạy xong và trả về (ok, mô tả)."""
    rc, out = remote_run(host, cmdline, interactive)
    if rc != 0:
        return False, '(%s) %s' % (rc, out)
    end = time.time() + timeout
    while time.time() < end:
        time.sleep(2)
        rc, out = run_args(['schtasks', '/Query', '/S', host, '/TN', 'KhoaVeyonRun', '/FO', 'CSV', '/V'], timeout=25)
        code = parse_task_result(out) if rc == 0 else None
        if code is not None and code not in (267009, 267011, 267014):  # đang chạy / chưa chạy / bị dừng
            return code == 0, 'xong' if code == 0 else 'mã lỗi %s' % code
    return False, 'quá thời gian chờ'


def job_cmd_lock(hosts, lock):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    folder = os.path.dirname(CONFIG_PATH)
    src = os.path.join(folder, CMDLOCK_NAME)
    with open(src, 'w', encoding='utf-8', newline='') as f:
        f.write(CMDLOCK_BAT.replace('\r\n', '\n').replace('\n', '\r\n'))
    target = cfg['dist_target'].rstrip('\\/')
    log('>>> %s %s (%d máy)' % ('Khóa' if lock else 'Mở khóa', target, len(hosts)), 'head')

    def one(h):
        try:
            distribute_to_host(h, [src], cfg['dist_target'])
        except Exception as e:  # noqa
            return False, 'không chép được %s: %s' % (CMDLOCK_NAME, e)
        ok, msg = remote_run_wait(h, 'cmd.exe /c %s\\%s %s <nul' % (target, CMDLOCK_NAME, 'lock' if lock else 'unlock'))
        return ok, ('đã khóa, học sinh chỉ đọc (log vẫn ghi được)' if lock else 'đã mở khóa') if ok else msg
    pool_run(one, hosts)


def job_setup_diagnose(hosts, push):
    cfg = json.loads(json.dumps(CFG))
    log('>>> Kiểm tra %d máy học sinh' % len(hosts), 'head')
    ready, results = [], []

    def one(h):
        checks = diagnose_host(cfg, h)
        return h, checks
    with ThreadPoolExecutor(max_workers=min(8, max(1, len(hosts)))) as ex:
        for h, checks in ex.map(one, hosts):
            results.append((h, checks))
            bad = [c for c in checks if not c[1]]
            log('[%s] %s' % (h, 'đã thiết lập đủ, không cần làm gì' if not bad else 'còn %d mục chưa đạt' % len(bad)), 'ok' if not bad else 'warn')
            for lab, ok, hint in checks:
                log('    %s %s%s' % ('✔' if ok else '✘', lab, '' if ok else '  →  ' + hint), 'ok' if ok else 'err')
            if not bad:
                continue
            if checks[1][1]:  # chép được file
                ready.append(h)
    need_install = [h for h, c in results if len(c) > 4 and not c[4][1]]
    if push and ready:
        files = [list_path(k) for k in LISTS if os.path.exists(list_path(k))]
        for f in sorted(os.listdir(PAYLOAD_DIR)) if os.path.isdir(PAYLOAD_DIR) else []:
            if f != STUDENT_SETUP_BAT:
                files.append(os.path.join(PAYLOAD_DIR, f))
        if files:
            job_dist(ready, files, 'none', '', '', False, cfg['dist_target'], 'Gửi file thiết lập xuống %d máy' % len(ready))
        todo = [h for h in ready if h in need_install]
        sname = os.path.basename((cfg.get('script_path') or '').replace('\\', '/'))
        if todo and os.path.isfile(os.path.join(PAYLOAD_DIR, sname)):
            log('>>> Cài dịch vụ (%s -install) trên %d máy' % (sname, len(todo)), 'head')

            def inst(h):
                # <nul để lệnh pause trong script không treo khi chạy nền bằng SYSTEM
                ok, msg = remote_run_wait(h, 'cmd.exe /c %s -install <nul' % cfg['script_path'])
                return ok, 'đã cài dịch vụ' if ok else '-install lỗi: ' + msg
            pool_run(inst, todo)
        elif todo:
            log('Chưa có %s trong payload nên chưa cài dịch vụ được.' % sname, 'warn')
        if cfg.get('auto_lock'):
            job_cmd_lock(ready, True)
    elif push:
        pending = [h for h, c in results if any(not x[1] for x in c)]
        if pending:
            log('Máy chưa chép được file (%s): chạy %s trên các máy này một lần (quyền quản trị), rồi bấm lại.' % (', '.join(pending), STUDENT_SETUP_BAT), 'warn')
        else:
            log('Tất cả máy đã thiết lập đủ.', 'ok')


def job_setup_all(hosts):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    log('>>> Thiết lập máy giáo viên', 'head')
    n = seed_default_lists(cfg)
    log('Danh sách Excel: đã tạo %d file mẫu' % n if n else 'Danh sách Excel: đã có đủ, giữ nguyên', 'ok')
    got, missing = ensure_payload(cfg)
    log('Đã tạo payload\\%s (chạy 1 lần trên mỗi máy học sinh)' % STUDENT_SETUP_BAT, 'ok')
    if got:
        log('Đã lấy %s vào payload' % ', '.join(got), 'ok')
    if missing:
        log('Chưa tìm thấy %s: hãy thêm vào payload (trang Phân phối & chạy)' % ', '.join(missing), 'warn')
    ok, msg = elevated_apply_veyon_items(cfg)
    log('Mục Veyon: ' + msg, 'ok' if ok else 'err')
    with CFG_LOCK:
        CFG['setup_done'] = True
        save_config(CFG)
    log('--- xong thiết lập máy giáo viên. Máy học sinh làm riêng bằng nút "Thiết lập máy học sinh" ---', 'head')


def run_apply_mode():
    """Chế độ phụ (chạy bằng quyền quản trị): ghi mục Veyon rồi lưu kết quả cho tiến trình chính đọc."""
    try:
        ok, msg = apply_veyon_items(CFG)
    except Exception as e:  # noqa
        ok, msg = False, 'Lỗi: %s' % e
    with open(SETUP_RESULT(), 'w', encoding='utf-8') as f:
        json.dump({'ok': ok, 'msg': msg}, f, ensure_ascii=False)


# ---------------------------------------------------------------- cấu hình Veyon: sao lưu, so sánh, khôi phục, đẩy xuống client, danh sách máy của Master
VCFG_KEY_RE = re.compile(r'^[A-Za-z0-9_]+(/[A-Za-z0-9_]+)+$')
VCFG_PRESETS = [  # (khóa, giá trị, mô tả). Chưa kiểm chứng trên Veyon thật: xem lại bằng "So sánh" sau khi đẩy.
    ('Network/PrimaryServicePort', '11100', 'Cổng Veyon Service'),
    ('Network/FirewallExceptionEnabled', 'true', 'Tự mở tường lửa cho Veyon'),
    ('Service/Autostart', 'true', 'Veyon Service tự khởi động cùng Windows'),
]
VCFG_RESTART = ['net stop VeyonService >nul 2>&1', 'net start VeyonService >nul 2>&1']


def vcfg_dir():
    d = os.path.join(os.path.dirname(CONFIG_PATH), 'veyon-backups')
    os.makedirs(d, exist_ok=True)
    return d


def vcfg_stamp():
    return time.strftime('%Y%m%d-%H%M%S')


def vcfg_safe(name):
    return re.sub(r'[^A-Za-z0-9.-]', '-', str(name))


def vcfg_remote_dir(cfg):
    return cfg['dist_target'].rstrip('\\/') + '\\vcfg'


def client_cli_path(cfg):
    return (cfg.get('client_veyon_cli') or r'C:\Program Files\Veyon\veyon-cli.exe').strip()


def vcfg_flatten(data):
    out = {}
    if not isinstance(data, dict):
        return out
    for g, sub in data.items():
        if isinstance(sub, dict):
            for k, v in sub.items():
                out['%s/%s' % (g, k)] = v
        else:
            out[str(g)] = sub
    return out


def vcfg_load_file(path):
    with open(path, 'r', encoding='utf-8-sig') as f:
        return vcfg_flatten(json.load(f))


def vcfg_diff(a, b):
    rows = []
    for k in sorted(set(a) | set(b)):
        if k not in b:
            rows.append(('-', k, a[k], None))
        elif k not in a:
            rows.append(('+', k, None, b[k]))
        elif a[k] != b[k]:
            rows.append(('~', k, a[k], b[k]))
    return rows


def _short(v, n=90):
    s = json.dumps(v, ensure_ascii=False)
    return s if len(s) <= n else s[:n] + '…'


def _bat(lines):
    return '@echo off\n' + '\n'.join(lines) + '\n'


def vcfg_export_master(cfg, dest):
    cli = find_veyon_cli(cfg['veyon_cli'])
    if not os.path.isfile(cli):
        return False, 'Chưa thấy veyon-cli.exe (đặt đường dẫn ở Cài đặt)'
    rc, out = run_args([cli, 'config', 'export', dest], timeout=60)
    if rc != 0 or not os.path.isfile(dest):
        return False, 'veyon-cli config export lỗi (%s): %s' % (rc, out)
    return True, ''


def client_script_run(host, cfg, bat_name, bat_text, files=None, timeout=90):
    """Chép file .bat (và các file kèm) vào <đích>\\vcfg trên máy học sinh, chạy bằng SYSTEM, đợi xong."""
    rdir = vcfg_remote_dir(cfg)
    tmp = tempfile.mkdtemp(prefix='kvcfg')
    try:
        bp = os.path.join(tmp, bat_name)
        with open(bp, 'w', encoding='utf-8', newline='') as f:
            f.write(bat_text.replace('\r\n', '\n').replace('\n', '\r\n'))
        srcs = [bp]
        for name, src in (files or {}).items():
            dst = os.path.join(tmp, name)
            shutil.copy2(src, dst)
            srcs.append(dst)
        distribute_to_host(host, srcs, rdir)
        return remote_run_wait(host, 'cmd.exe /c %s\\%s <nul' % (rdir, bat_name), timeout)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def client_fetch(host, cfg, name, dest):
    src = os.path.join(unc_target(host, vcfg_remote_dir(cfg)), name)
    if not os.path.isfile(src):
        return False
    shutil.copy2(src, dest)
    try:
        os.remove(src)
    except OSError:
        pass
    return True


def vcfg_export_client(cfg, host, dest):
    t, cli = vcfg_remote_dir(cfg), client_cli_path(cfg)
    ok, msg = client_script_run(host, cfg, 'vcfg-export.bat', _bat([
        '"%s" config export "%s\\export.json"' % (cli, t), 'exit /b %errorlevel%']))
    if not ok:
        return False, msg
    if not client_fetch(host, cfg, 'export.json', dest):
        return False, 'chạy xong nhưng không thấy file xuất ra (kiểm tra đường dẫn veyon-cli trên máy học sinh ở khối Cài đặt phía dưới)'
    return True, ''


def vcfg_backup_client(cfg, host):
    dest = os.path.join(vcfg_dir(), '%s_%s.json' % (vcfg_safe(host), vcfg_stamp()))
    ok, msg = vcfg_export_client(cfg, host, dest)
    return ok, ('đã lưu ' + os.path.basename(dest)) if ok else msg


def vcfg_restore_client(cfg, host, path):
    t, cli = vcfg_remote_dir(cfg), client_cli_path(cfg)
    ok, msg = client_script_run(host, cfg, 'vcfg-restore.bat', _bat(
        ['"%s" config export "%s\\before-restore.json"' % (cli, t),
         '"%s" config import "%s\\restore.json"' % (cli, t), 'set R=%errorlevel%'] + VCFG_RESTART + ['exit /b %R%']),
        files={'restore.json': path}, timeout=120)
    got = client_fetch(host, cfg, 'before-restore.json',
                       os.path.join(vcfg_dir(), '%s_truoc-khoi-phuc_%s.json' % (vcfg_safe(host), vcfg_stamp())))
    note = ' (đã sao lưu cấu hình cũ của máy)' if got else ''
    try:
        os.remove(os.path.join(unc_target(host, t), 'restore.json'))
    except OSError:
        pass
    return ok, ('đã khôi phục' + note) if ok else '%s%s' % (msg, note)


def vcfg_check_items(items):
    out = []
    for it in items or []:
        k, v = str(it.get('key', '')).strip(), str(it.get('value', ''))
        if not k and not v.strip():
            continue
        if not VCFG_KEY_RE.match(k):
            raise ApiError('Khóa "%s" không hợp lệ, phải dạng Nhóm/Tên (ví dụ Network/PrimaryServicePort)' % k)
        if '"' in v or '\n' in v or '\r' in v:
            raise ApiError('Giá trị của %s không được chứa dấu " hoặc xuống dòng' % k)
        out.append((k, v))
    if not out:
        raise ApiError('Chưa có khóa cấu hình nào để đẩy')
    return out


def vcfg_push_client(cfg, host, items, restart):
    t, cli = vcfg_remote_dir(cfg), client_cli_path(cfg)
    lines = ['chcp 65001 >nul', '"%s" config export "%s\\before-push.json"' % (cli, t), 'set F=0']
    for k, v in items:
        lines.append('"%s" config set "%s" "%s"' % (cli, k, v.replace('%', '%%')))
        lines.append('if errorlevel 1 set F=1')
    if restart:
        lines += VCFG_RESTART
    lines.append('exit /b %F%')
    ok, msg = client_script_run(host, cfg, 'vcfg-push.bat', _bat(lines), timeout=120)
    got = client_fetch(host, cfg, 'before-push.json',
                       os.path.join(vcfg_dir(), '%s_truoc-khi-day_%s.json' % (vcfg_safe(host), vcfg_stamp())))
    note = ' (đã sao lưu cấu hình cũ)' if got else ''
    return ok, ('đã đẩy %d khóa%s%s' % (len(items), ', đã khởi động lại Veyon Service' if restart else '', note)) if ok else (
        'có khóa bị lỗi hoặc không chạy được: %s%s' % (msg, note))


def vcfg_restore_master(cfg, path):
    cli = find_veyon_cli(cfg['veyon_cli'])
    if not os.path.isfile(cli):
        return False, 'Chưa thấy veyon-cli.exe'
    ok, msg = vcfg_export_master(cfg, os.path.join(vcfg_dir(), 'master_truoc-khoi-phuc_%s.json' % vcfg_stamp()))
    if not ok:
        return False, 'Không sao lưu được cấu hình hiện tại nên dừng: ' + msg
    rc, out = run_args([cli, 'config', 'import', path], timeout=60)
    if rc != 0:
        return False, 'veyon-cli config import lỗi (%s): %s' % (rc, out)
    return True, 'Đã khôi phục cấu hình master (cấu hình cũ đã lưu trong veyon-backups). Hãy mở lại Veyon Master.'


def vcfg_netobj_apply(cfg, hosts, location):
    """Thêm máy vào danh sách máy của Veyon Master bằng veyon-cli networkobject. Máy đã có thì bỏ qua."""
    cli = find_veyon_cli(cfg['veyon_cli'])
    if not os.path.isfile(cli):
        return False, 'Chưa thấy veyon-cli.exe'
    rc, out = run_args([cli, 'networkobject', 'list'], timeout=30)
    existing = out.lower() if rc == 0 else ''
    location = (location or '').strip()
    if location and location.lower() not in existing:
        rc, out = run_args([cli, 'networkobject', 'add', 'location', location], timeout=30)
        if rc != 0:
            return False, 'Không tạo được nhóm "%s" (%s): %s' % (location, rc, out)
    added = skipped = 0
    errs = []
    for h in hosts:
        if re.search(r'(?<![\w.-])%s(?![\w.-])' % re.escape(h.lower()), existing):
            skipped += 1
            continue
        args = [cli, 'networkobject', 'add', 'computer', h, h, '']
        if location:
            args.append(location)
        rc, out = run_args(args, timeout=30)
        if rc == 0:
            added += 1
        else:
            errs.append('%s (%s) %s' % (h, rc, out))
    msg = 'Đã thêm %d máy, bỏ qua %d máy đã có trong danh sách.' % (added, skipped)
    if errs:
        return False, msg + ' Lỗi: ' + '; '.join(errs[:5])
    return True, msg + ' Hãy mở lại Veyon Master để thấy danh sách mới.'


ELEV_OPS = {
    'netobj': lambda cfg, a: vcfg_netobj_apply(cfg, a.get('hosts') or [], a.get('location', '')),
    'restore': lambda cfg, a: vcfg_restore_master(cfg, a['path']),
}
ELEV_JOB = lambda: os.path.join(os.path.dirname(CONFIG_PATH), 'elevated-job.json')  # noqa


def run_elevated_job():
    """Chế độ phụ chạy bằng quyền quản trị: làm một việc trong ELEV_OPS rồi ghi kết quả cho tiến trình chính."""
    try:
        with open(ELEV_JOB(), 'r', encoding='utf-8') as f:
            job = json.load(f)
        ok, msg = ELEV_OPS[job['op']](CFG, job.get('args') or {})
    except Exception as e:  # noqa
        ok, msg = False, 'Lỗi: %s' % e
    with open(SETUP_RESULT(), 'w', encoding='utf-8') as f:
        json.dump({'ok': ok, 'msg': msg}, f, ensure_ascii=False)
    try:
        os.remove(ELEV_JOB())
    except OSError:
        pass


def elevated_op(op, args):
    cfg = json.loads(json.dumps(CFG))
    if is_admin():
        return ELEV_OPS[op](cfg, args)
    with open(ELEV_JOB(), 'w', encoding='utf-8') as f:
        json.dump({'op': op, 'args': args}, f, ensure_ascii=False)
    try:
        os.remove(SETUP_RESULT())
    except OSError:
        pass
    import ctypes
    params = '"%s" --elevated-job' % os.path.abspath(__file__)
    r = ctypes.windll.shell32.ShellExecuteW(None, 'runas', sys.executable, params, APP_DIR, 0)  # noqa
    if r <= 32:
        return False, 'Bạn đã từ chối quyền quản trị (hoặc không mở được UAC)'
    for _ in range(120):
        time.sleep(0.5)
        if os.path.exists(SETUP_RESULT()):
            time.sleep(0.2)
            try:
                with open(SETUP_RESULT(), 'r', encoding='utf-8') as f:
                    d = json.load(f)
                return bool(d.get('ok')), d.get('msg', '')
            except Exception:  # noqa
                continue
    return False, 'Quá thời gian chờ kết quả từ cửa sổ quản trị'


def vcfg_snapshot(cfg, ref, tmpdir):
    """ref: 'master' | 'host:TÊN' (đọc trực tiếp) | tên file trong veyon-backups. Trả về (dict phẳng, lỗi)."""
    if ref == 'master':
        p = os.path.join(tmpdir, 'master.json')
        ok, msg = vcfg_export_master(cfg, p)
    elif ref.startswith('host:'):
        p = os.path.join(tmpdir, vcfg_safe(ref[5:]) + '.json')
        ok, msg = vcfg_export_client(cfg, ref[5:], p)
    else:
        p, ok, msg = vcfg_backup_path(ref), True, ''
    if not ok:
        return None, msg
    try:
        return vcfg_load_file(p), ''
    except Exception as e:  # noqa
        return None, 'Không đọc được %s: %s' % (ref, e)


def vcfg_backup_path(name):
    name = os.path.basename(str(name or ''))
    p = os.path.join(vcfg_dir(), name)
    if not name.lower().endswith('.json') or not os.path.isfile(p):
        raise ApiError('Không thấy file sao lưu "%s"' % name)
    return p


def job_vcfg_backup(hosts, master):
    cfg = json.loads(json.dumps(CFG))
    log('>>> Sao lưu cấu hình Veyon (%s%s)' % ('master' if master else '', (' + ' if master and hosts else '') + ('%d máy học sinh' % len(hosts) if hosts else '')), 'head')
    if master:
        dest = os.path.join(vcfg_dir(), 'master_%s.json' % vcfg_stamp())
        ok, msg = vcfg_export_master(cfg, dest)
        log('[master] %s' % (('OK đã lưu ' + os.path.basename(dest)) if ok else 'LỖI ' + msg), 'ok' if ok else 'err')
    if hosts:
        pool_run(lambda h: vcfg_backup_client(cfg, h), hosts)


def job_vcfg_restore(target, hosts, path):
    cfg = json.loads(json.dumps(CFG))
    log('>>> Khôi phục cấu hình Veyon từ %s' % os.path.basename(path), 'head')
    if target == 'master':
        ok, msg = elevated_op('restore', {'path': path})
        log('[master] %s %s' % ('OK' if ok else 'LỖI', msg), 'ok' if ok else 'err')
    else:
        pool_run(lambda h: vcfg_restore_client(cfg, h, path), hosts)


def job_vcfg_push(hosts, items, restart):
    cfg = json.loads(json.dumps(CFG))
    log('>>> Đẩy %d khóa cấu hình Veyon xuống %d máy' % (len(items), len(hosts)), 'head')
    for k, v in items:
        log('    %s = %s' % (k, v))
    pool_run(lambda h: vcfg_push_client(cfg, h, items, restart), hosts)


def job_vcfg_netobj(hosts, location):
    log('>>> Thêm %d máy vào danh sách máy của Veyon Master' % len(hosts), 'head')
    ok, msg = elevated_op('netobj', {'hosts': hosts, 'location': location})
    log('[master] %s %s' % ('OK' if ok else 'LỖI', msg), 'ok' if ok else 'err')


def job_vcfg_compare(a, b):
    cfg = json.loads(json.dumps(CFG))
    log('>>> So sánh cấu hình Veyon: %s  ↔  %s' % (a, b), 'head')
    tmp = tempfile.mkdtemp(prefix='kvcmp')
    try:
        da, ea = vcfg_snapshot(cfg, a, tmp)
        if da is None:
            log('LỖI bên A: ' + ea, 'err')
            return
        db, eb = vcfg_snapshot(cfg, b, tmp)
        if db is None:
            log('LỖI bên B: ' + eb, 'err')
            return
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    rows = vcfg_diff(da, db)
    if not rows:
        log('Giống nhau hoàn toàn (%d khóa).' % len(da), 'ok')
        return
    for sign, k, va, vb in rows[:300]:
        if sign == '~':
            log('  ~ %s:  A = %s   |   B = %s' % (k, _short(va), _short(vb)), 'warn')
        elif sign == '+':
            log('  + %s chỉ có ở B = %s' % (k, _short(vb)), 'ok')
        else:
            log('  - %s chỉ có ở A = %s' % (k, _short(va)), 'err')
    log('--- %d khác biệt%s ---' % (len(rows), ' (chỉ hiện 300 dòng đầu)' if len(rows) > 300 else ''), 'head')


def api_vcfg_info(_q=None):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    files = []
    for f in sorted(os.listdir(vcfg_dir()), reverse=True):
        p = os.path.join(vcfg_dir(), f)
        if f.lower().endswith('.json') and os.path.isfile(p):
            files.append({'name': f, 'size': os.path.getsize(p), 'time': time.strftime('%d/%m %H:%M', time.localtime(os.path.getmtime(p)))})
    return {'backups': files, 'dir': vcfg_dir(), 'push': cfg.get('vcfg_push') or [], 'restart': bool(cfg.get('vcfg_restart', True)),
            'location': cfg.get('netobj_location', ''), 'client_cli': client_cli_path(cfg),
            'presets': [{'key': k, 'value': v, 'desc': d} for k, v, d in VCFG_PRESETS]}


def api_vcfg_save(b):
    with CFG_LOCK:
        push = []
        for it in (b.get('push') or [])[:60]:
            push.append({'key': str(it.get('key', '')).strip(), 'value': str(it.get('value', ''))})
        CFG['vcfg_push'] = push
        CFG['vcfg_restart'] = bool(b.get('restart', True))
        CFG['netobj_location'] = str(b.get('location', '')).strip()
        CFG['client_veyon_cli'] = str(b.get('client_cli', '')).strip() or DEFAULT_CONFIG['client_veyon_cli']
        save_config(CFG)
    return {'ok': True}


def api_vcfg_run(b):
    step = b.get('step')
    hosts = clean_hosts(b.get('hosts'))
    if step in ('backup_master', 'backup_clients', 'backup_all'):
        if step != 'backup_master' and not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        start_job(job_vcfg_backup, hosts if step != 'backup_master' else [], step != 'backup_clients')
    elif step == 'compare':
        a, c = str(b.get('a', '')), str(b.get('b', ''))
        if not a or not c or a == c:
            raise ApiError('Hãy chọn hai bên khác nhau để so sánh')
        for ref in (a, c):
            if ref != 'master' and not ref.startswith('host:'):
                vcfg_backup_path(ref)
        start_job(job_vcfg_compare, a, c)
    elif step in ('restore_master', 'restore_clients'):
        path = vcfg_backup_path(b.get('file'))
        if step == 'restore_clients' and not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        start_job(job_vcfg_restore, 'master' if step == 'restore_master' else 'clients', hosts, path)
    elif step == 'push':
        if not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        items = vcfg_check_items(b.get('push'))
        start_job(job_vcfg_push, hosts, items, bool(b.get('restart', True)))
    elif step == 'netobj':
        if not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        start_job(job_vcfg_netobj, hosts, str(b.get('location', '')).strip())
    else:
        raise ApiError('Bước không hợp lệ')
    return {'ok': True}


# ---------------------------------------------------------------- chụp màn hình máy học sinh (trước khi khóa)
SHOT_PS1 = r"""param([string]$Out)
Add-Type -AssemblyName System.Windows.Forms,System.Drawing
Add-Type -TypeDefinition 'using System.Runtime.InteropServices; public class KvDpi { [DllImport("user32.dll")] public static extern bool SetProcessDPIAware(); }'
[KvDpi]::SetProcessDPIAware() | Out-Null
$b = [System.Windows.Forms.SystemInformation]::VirtualScreen
$bmp = New-Object System.Drawing.Bitmap $b.Width, $b.Height
$g = [System.Drawing.Graphics]::FromImage($bmp)
$g.CopyFromScreen($b.Left, $b.Top, 0, 0, $bmp.Size)
$bmp.Save($Out, [System.Drawing.Imaging.ImageFormat]::Jpeg)
$g.Dispose(); $bmp.Dispose()
"""
SHOT_REMOTE_FILE = r'C:\Users\Public\kvshot.jpg'  # thư mục mà mọi người dùng đều ghi được
SHOT_PS1_NAME = 'kvshot.ps1'


def shot_root():
    d = os.path.join(os.path.dirname(CONFIG_PATH), 'screenshots')
    os.makedirs(d, exist_ok=True)
    return d


def _q(p):
    return '"%s"' % p if ' ' in p else p


def shot_one(cfg, host, folder):
    """Chụp màn hình người đang đăng nhập trên máy học sinh rồi lấy ảnh về máy giáo viên. Trả về (ok, mô tả)."""
    rdir = vcfg_remote_dir(cfg)
    tmp = tempfile.mkdtemp(prefix='kvshot')
    try:
        sp = os.path.join(tmp, SHOT_PS1_NAME)
        with open(sp, 'w', encoding='utf-8-sig', newline='') as f:  # BOM để PowerShell 5 đọc đúng
            f.write(SHOT_PS1.replace('\r\n', '\n').replace('\n', '\r\n'))
        distribute_to_host(host, [sp], rdir)
    except Exception as e:  # noqa
        return False, 'không chép được script chụp: %s' % e
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    import ntpath
    remote_img = os.path.join(unc_target(host, ntpath.dirname(SHOT_REMOTE_FILE)), ntpath.basename(SHOT_REMOTE_FILE))
    try:
        os.remove(remote_img)
    except OSError:
        pass
    cmd = 'powershell.exe -NoProfile -ExecutionPolicy Bypass -WindowStyle Hidden -File %s -Out %s' % (
        _q(rdir + '\\' + SHOT_PS1_NAME), _q(SHOT_REMOTE_FILE))
    ok, msg = remote_run_wait(host, cmd, 45, interactive=True)
    if not ok:
        return False, msg
    dest = os.path.join(folder, vcfg_safe(host) + '.jpg')
    try:
        shutil.copy2(remote_img, dest)
        os.remove(remote_img)
    except OSError as e:
        return False, 'chụp xong nhưng không lấy được ảnh: %s' % e
    return True, 'đã lưu ' + os.path.basename(dest)


def job_shot(hosts, title='Chụp màn hình'):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    folder = os.path.join(shot_root(), time.strftime('%Y%m%d-%H%M%S'))
    os.makedirs(folder, exist_ok=True)
    log('>>> %s (%d máy) -> %s' % (title, len(hosts), folder), 'head')
    pool_run(lambda h: shot_one(cfg, h, folder), hosts)


def job_shot_then_action(key, hosts):
    job_shot(hosts, 'Chụp màn hình trước khi "%s"' % ACTION_BY_KEY[key][1])
    job_action(key, hosts)  # máy chụp lỗi vẫn bị khóa, xem lỗi ở Nhật ký


def api_shot(b):
    hosts = clean_hosts(b.get('hosts'))
    if not hosts:
        raise ApiError('Chưa có máy nào để chụp')
    start_job(job_shot, hosts)
    return {'ok': True}


# ---------------------------------------------------------------- giám sát vi phạm: tự chụp ảnh khi học sinh mở trang/ứng dụng ngoài danh sách
WATCH_REMOTE_DIR = r'C:\Users\Public\KhoaVeyonEvidence'  # thư mục mà mọi người dùng đều ghi được
WATCH_PS1_NAME = 'KhoaVeyonWatch.ps1'
WATCH_TASK = 'KhoaVeyonWatch'
WATCH_PS1 = r'''param(
    [string]$ListDir = 'C:\cmd',
    [string]$Out = 'C:\Users\Public\KhoaVeyonEvidence',
    [int]$RepeatSec = 10,
    [int]$MaxShots = 5
)
# KhoaVeyonWatch: chay trong phien dang nhap cua hoc sinh. Khi cua so tren cung la trang web khong co trong
# AllowPage.xlsx hoac ung dung khong co trong KeepAppList.xlsx thi chup man hinh (khong ghi phim go, khong ghi noi dung go).
$ErrorActionPreference = 'SilentlyContinue'
$created = $false
$mx = New-Object System.Threading.Mutex($true, 'Local\KhoaVeyonWatch', [ref]$created)
if (-not $created) { exit }
New-Item -ItemType Directory -Path $Out -Force | Out-Null
$stopFile = Join-Path $Out 'stop.flag'
Remove-Item -LiteralPath $stopFile -Force
$pc = $env:COMPUTERNAME; $user = $env:USERNAME
$csv = Join-Path $Out 'evidence.csv'
$wlog = Join-Path $Out 'watch.log'
function Log($m) { try { Add-Content -LiteralPath $wlog -Value ((Get-Date -Format 'yyyy-MM-dd HH:mm:ss') + '  ' + $m) } catch {} }
Add-Type -AssemblyName System.Windows.Forms, System.Drawing, UIAutomationClient, UIAutomationTypes
Add-Type -TypeDefinition @'
using System;
using System.Runtime.InteropServices;
public class KvW {
    [DllImport("user32.dll")] public static extern IntPtr GetForegroundWindow();
    [DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr h, out uint pid);
    [DllImport("user32.dll")] public static extern int GetWindowText(IntPtr h, System.Text.StringBuilder s, int n);
    [DllImport("user32.dll")] public static extern bool SetProcessDPIAware();
}
'@
[void][KvW]::SetProcessDPIAware()

function Read-ColumnA($file) {
    $out = New-Object System.Collections.Generic.List[string]
    if (-not (Test-Path -LiteralPath $file)) { Log "File not found: $file"; return @() }
    Add-Type -AssemblyName System.IO.Compression
    $fs  = [IO.File]::Open($file, [IO.FileMode]::Open, [IO.FileAccess]::Read, [IO.FileShare]::ReadWrite)
    $zip = New-Object IO.Compression.ZipArchive($fs, [IO.Compression.ZipArchiveMode]::Read)
    try {
        $shared = @()
        $e = $zip.GetEntry('xl/sharedStrings.xml')
        if ($e) {
            $r = New-Object IO.StreamReader($e.Open())
            $x = New-Object Xml.XmlDocument
            $x.LoadXml($r.ReadToEnd()); $r.Close()
            foreach ($si in $x.SelectNodes('//*[local-name()="si"]')) {
                $shared += (($si.SelectNodes('.//*[local-name()="t"]') | ForEach-Object { $_.InnerText }) -join '')
            }
        }
        $e = $zip.GetEntry('xl/worksheets/sheet1.xml')
        if (-not $e) {
            $e = $zip.Entries | Where-Object { $_.FullName -like 'xl/worksheets/sheet*.xml' } | Sort-Object FullName | Select-Object -First 1
        }
        if (-not $e) { Log "No worksheet found in $file"; return @() }
        $r = New-Object IO.StreamReader($e.Open())
        $x = New-Object Xml.XmlDocument
        $x.LoadXml($r.ReadToEnd()); $r.Close()
        foreach ($c in $x.SelectNodes('//*[local-name()="c"]')) {
            if ($c.GetAttribute('r') -notmatch '^A\d+$') { continue }
            $t = $c.GetAttribute('t'); $val = $null
            if ($t -eq 's') {
                $v = $c.SelectSingleNode('*[local-name()="v"]')
                if ($v) { $val = $shared[[int]$v.InnerText] }
            } elseif ($t -eq 'inlineStr') {
                $val = ($c.SelectNodes('.//*[local-name()="t"]') | ForEach-Object { $_.InnerText }) -join ''
            } else {
                $v = $c.SelectSingleNode('*[local-name()="v"]')
                if ($v) { $val = $v.InnerText }
            }
            if ($val) { $out.Add($val.Trim()) }
        }
    } finally { $zip.Dispose(); $fs.Dispose() }
    return $out.ToArray()
}

$browsers = @('chrome','msedge','firefox','brave','opera','coccoc','vivaldi')
$protect = @('explorer','cmd','conhost','powershell','pwsh','veyon-server','veyon-worker','veyon-service','veyon-master',
             'applicationframehost','shellexperiencehost','startmenuexperiencehost','searchhost','textinputhost',
             'systemsettings','sihost','ctfmon','dwm','lockapp','searchapp')
$allowDomains = @(); $keepNames = @(); $keepPaths = @()
function Load-Lists {
    $d1 = @(); $n1 = @(); $p1 = @()
    foreach ($i in @(Read-ColumnA (Join-Path $ListDir 'AllowPage.xlsx'))) {
        $d = (($i.ToLower() -replace '^[a-z]+://', '') -replace '[/?#:].*$', '') -replace '^www\.', ''
        if ($d -match '^[a-z0-9-]+(\.[a-z0-9-]+)+$') { $d1 += $d }
    }
    foreach ($k in @(Read-ColumnA (Join-Path $ListDir 'KeepAppList.xlsx'))) {
        $k = [Environment]::ExpandEnvironmentVariables($k).Trim('"')
        if ($k -match '\\') { $p1 += $k.TrimEnd('\') }
        elseif ($k) { $n1 += ($k -replace '\.exe$', '').ToLower() }
    }
    $script:allowDomains = $d1; $script:keepNames = $n1; $script:keepPaths = $p1
}
function Site-Allowed($h) {
    foreach ($d in $allowDomains) { if ($h -eq $d -or $h.EndsWith('.' + $d)) { return $true } }
    return $false
}
function App-Allowed($n, $path) {
    if ($keepNames -contains $n) { return $true }
    if ($path) { foreach ($kp in $keepPaths) { if ($path -like "$kp*") { return $true } } }
    return $false
}
# Tra ve ten mien dang mo, '' neu la trang noi bo (newtab...), $null neu khong doc duoc thanh dia chi
function Get-BrowserHost($hwnd) {
    try {
        $root = [System.Windows.Automation.AutomationElement]::FromHandle($hwnd)
        $cond = New-Object System.Windows.Automation.PropertyCondition([System.Windows.Automation.AutomationElement]::ControlTypeProperty, [System.Windows.Automation.ControlType]::Edit)
        foreach ($e in $root.FindAll([System.Windows.Automation.TreeScope]::Descendants, $cond)) {
            $vp = $null
            if ($e.TryGetCurrentPattern([System.Windows.Automation.ValuePattern]::Pattern, [ref]$vp)) {
                $v = ([string]$vp.Current.Value).Trim().ToLower()
                if ($v -match '^(chrome|edge|about|brave|opera|vivaldi)://') { return '' }
                if ($v -match '^(https?://)?([a-z0-9-]+\.)+[a-z0-9-]{2,}(:\d+)?([/?#].*)?$') {
                    return (($v -replace '^[a-z]+://', '') -replace '[/?#:].*$', '') -replace '^www\.', ''
                }
            }
        }
    } catch {}
    return $null
}
function Take-Shot($kind, $name, $title) {
    try {
        $now = Get-Date
        $safe = ($name -replace '[^a-z0-9.-]', '_')
        $file = Join-Path $Out ('{0}_{1}_{2}_{3}.jpg' -f $pc, $now.ToString('yyyyMMdd-HHmmss'), $kind, $safe)
        $b = [System.Windows.Forms.SystemInformation]::VirtualScreen
        $bmp = New-Object System.Drawing.Bitmap $b.Width, $b.Height
        $g = [System.Drawing.Graphics]::FromImage($bmp)
        $g.CopyFromScreen($b.Left, $b.Top, 0, 0, $bmp.Size)
        $bmp.Save($file, [System.Drawing.Imaging.ImageFormat]::Jpeg)
        $g.Dispose(); $bmp.Dispose()
        if (-not (Test-Path -LiteralPath $csv)) { Set-Content -LiteralPath $csv -Value 'Date,Time,PCName,User,Type,Name,Window,File' -Encoding UTF8 }
        $t = $title -replace '"', '""'
        Add-Content -LiteralPath $csv -Encoding UTF8 -Value ('"{0}","{1}","{2}","{3}","{4}","{5}","{6}","{7}"' -f $now.ToString('yyyy-MM-dd'), $now.ToString('HH:mm:ss'), $pc, $user, $kind, $name, $t, (Split-Path $file -Leaf))
    } catch { Log "Shot failed: $_" }
}

Load-Lists
Log ("Watch STARTED. Sites allowed: {0}; apps allowed: {1} names, {2} paths" -f $allowDomains.Count, $keepNames.Count, $keepPaths.Count)
$lastLoad = Get-Date; $curKey = ''; $shots = 0; $lastShot = [datetime]::MinValue
while (-not (Test-Path -LiteralPath $stopFile)) {
    try {
        if (((Get-Date) - $lastLoad).TotalSeconds -ge 30) { Load-Lists; $lastLoad = Get-Date }
        $h = [KvW]::GetForegroundWindow()
        $sb = New-Object System.Text.StringBuilder 256
        [void][KvW]::GetWindowText($h, $sb, 256)
        $title = $sb.ToString()
        $procId = 0
        [void][KvW]::GetWindowThreadProcessId($h, [ref]$procId)
        $p = Get-Process -Id $procId -ErrorAction SilentlyContinue
        $n = ''; $path = $null
        if ($p) { $n = $p.ProcessName.ToLower(); try { $path = $p.Path } catch {} }
        $kind = ''; $name = ''
        if ($n -and ($protect -notcontains $n) -and ($n -ne 'idle')) {
            if ($browsers -contains $n) {
                $site = Get-BrowserHost $h
                if ($site -and $allowDomains.Count -gt 0 -and -not (Site-Allowed $site)) { $kind = 'site'; $name = $site }
            } elseif (($keepNames.Count + $keepPaths.Count) -gt 0 -and -not (App-Allowed $n $path)) {
                $kind = 'app'; $name = $n
            }
        }
        if ($kind) {
            $key = "$kind|$name"
            if ($key -ne $curKey) { $curKey = $key; $shots = 0; $lastShot = [datetime]::MinValue }
            if ($shots -lt $MaxShots -and ((Get-Date) - $lastShot).TotalSeconds -ge $RepeatSec) {
                Take-Shot $kind $name $title; $shots++; $lastShot = Get-Date
            }
        } else { $curKey = '' }
    } catch { Log "Loop error: $_" }
    Start-Sleep -Seconds 2
}
Log 'Watch STOPPED'
'''


def watch_num(cfg, key, default, lo, hi):
    try:
        return max(lo, min(hi, int(str(cfg.get(key) or default).strip())))
    except ValueError:
        return default


def evidence_root():
    d = os.path.join(os.path.dirname(CONFIG_PATH), 'evidence')
    os.makedirs(d, exist_ok=True)
    return d


def watch_start_one(cfg, host):
    rdir = vcfg_remote_dir(cfg)
    tmp = tempfile.mkdtemp(prefix='kvwatch')
    try:
        sp = os.path.join(tmp, WATCH_PS1_NAME)
        with open(sp, 'w', encoding='utf-8-sig', newline='') as f:
            f.write(WATCH_PS1.replace('\r\n', '\n').replace('\n', '\r\n'))
        distribute_to_host(host, [sp], rdir)
    except Exception as e:  # noqa
        return False, 'không chép được script giám sát: %s' % e
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    ev = unc_target(host, WATCH_REMOTE_DIR)
    try:
        os.makedirs(ev, exist_ok=True)
        os.remove(os.path.join(ev, 'stop.flag'))
    except OSError:
        pass
    cmd = 'powershell.exe -NoProfile -ExecutionPolicy Bypass -WindowStyle Hidden -File %s -ListDir %s -Out %s -RepeatSec %d -MaxShots %d' % (
        _q(rdir + '\\' + WATCH_PS1_NAME), _q(cfg['dist_target'].rstrip('\\/')), _q(WATCH_REMOTE_DIR),
        watch_num(cfg, 'watch_repeat', 10, 3, 300), watch_num(cfg, 'watch_max', 5, 1, 50))
    rc, out = remote_run(host, cmd, interactive=True, name=WATCH_TASK)
    return rc == 0, 'đã bật giám sát' if rc == 0 else out


def watch_stop_one(cfg, host):
    ev = unc_target(host, WATCH_REMOTE_DIR)
    os.makedirs(ev, exist_ok=True)
    with open(os.path.join(ev, 'stop.flag'), 'w') as f:
        f.write('stop')
    return True, 'đã gửi lệnh dừng (dừng sau vài giây)'


def watch_fetch_one(cfg, host):
    src = unc_target(host, WATCH_REMOTE_DIR)
    if not os.path.isdir(src):
        return True, 'chưa có ảnh vi phạm'
    dest = os.path.join(evidence_root(), vcfg_safe(host))
    os.makedirs(dest, exist_ok=True)
    n = 0
    for f in os.listdir(src):
        p = os.path.join(src, f)
        low = f.lower()
        if low.endswith('.jpg'):
            shutil.copy2(p, os.path.join(dest, f))
            os.remove(p)  # đã có bản ở máy giáo viên, xóa bản trên máy học sinh
            n += 1
        elif low in ('evidence.csv', 'watch.log'):
            shutil.copy2(p, os.path.join(dest, f))
    return True, 'đã lấy %d ảnh về %s' % (n, dest)


def job_watch(step, hosts):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    fn = {'start': watch_start_one, 'stop': watch_stop_one, 'fetch': watch_fetch_one}[step]
    log('>>> %s giám sát vi phạm (%d máy)' % ({'start': 'Bật', 'stop': 'Tắt', 'fetch': 'Lấy ảnh'}[step], len(hosts)), 'head')
    pool_run(lambda h: fn(cfg, h), hosts)


def api_watch(b):
    step = b.get('step')
    hosts = clean_hosts(b.get('hosts'))
    if step not in ('start', 'stop', 'fetch'):
        raise ApiError('Bước không hợp lệ')
    if not hosts:
        raise ApiError('Chưa có máy nào')
    start_job(job_watch, step, hosts)
    return {'ok': True}


# ---------------------------------------------------------------- tự cập nhật từ GitHub
import zipfile
import tempfile

UPDATE_FILES = ('KhoaVeyon.py', 'Chay-KhoaVeyon.bat', 'README-KhoaVeyon.md', 'requirements.txt')
UPDATE_DIRS = ('web',)          # chỉ ghi đè mã chương trình; data, lists, payload, runtime của bạn không bị đụng tới
UPDATE_MAX = 60 * 1024 * 1024
UPDATE_HOSTS = ('github.com', 'githubusercontent.com')


def _ver(text):
    nums = re.findall(r'\d+', str(text))
    return tuple(int(x) for x in nums[:4]) or (0,)


def _host_ok(url):
    h = (urllib.parse.urlparse(url).hostname or '').lower()
    return urllib.parse.urlparse(url).scheme == 'https' and any(h == d or h.endswith('.' + d) for d in UPDATE_HOSTS)


def _http_get(url, timeout=15, limit=UPDATE_MAX, trusted=False):
    if not trusted and not _host_ok(url):
        raise ApiError('Địa chỉ cập nhật không hợp lệ (chỉ chấp nhận https của GitHub)')
    req = urllib.request.Request(url, headers={'User-Agent': 'KhoaVeyon/' + VERSION, 'Accept': 'application/vnd.github+json'})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        data = r.read(limit + 1)
    if len(data) > limit:
        raise ApiError('File cập nhật quá lớn')
    return data


def check_update(repo, api_base='https://api.github.com', trusted=False):
    """Hỏi GitHub Releases xem có bản mới hơn VERSION không. Không tải gì về."""
    if not re.fullmatch(r'[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+', repo or ''):
        raise ApiError('Chưa điền kho GitHub (dạng chủ-sở-hữu/tên-kho) ở trang Cài đặt')
    try:
        data = json.loads(_http_get('%s/repos/%s/releases/latest' % (api_base, repo), trusted=trusted).decode('utf-8'))
    except ApiError:
        raise
    except urllib.error.HTTPError as e:
        if e.code == 404:
            raise ApiError('Kho %s chưa có Release nào (hoặc là kho riêng tư). Tạo Release trên GitHub rồi kiểm tra lại.' % repo)
        if e.code in (403, 429):
            raise ApiError('GitHub tạm giới hạn số lần hỏi, hãy thử lại sau ít phút.')
        raise ApiError('GitHub trả lỗi %s khi kiểm tra bản mới' % e.code)
    except Exception as e:  # noqa
        raise ApiError('Không kiểm tra được bản mới (cần Internet, và kho phải công khai có Release): %s' % e)
    tag = str(data.get('tag_name') or '')
    url = ''
    for a in data.get('assets') or []:
        n = str(a.get('name', ''))
        if n.lower().endswith('.zip') and 'khoaveyon' in n.lower():
            url = a.get('browser_download_url', '')
            break
    url = url or data.get('zipball_url', '')
    return {'current': VERSION, 'latest': tag, 'newer': _ver(tag) > _ver(VERSION), 'url': url,
            'notes': str(data.get('body') or '')[:1500], 'page': data.get('html_url', '')}


def apply_update(url, app_dir=None, trusted=False):
    """Tải bản mới, chỉ thay file chương trình (có sao lưu bản cũ). Trả về danh sách file đã thay."""
    app_dir = app_dir or APP_DIR
    blob = _http_get(url, timeout=60, trusted=trusted)
    tmp = tempfile.mkdtemp(prefix='khoaveyon-upd-')
    try:
        zp = os.path.join(tmp, 'u.zip')
        with open(zp, 'wb') as f:
            f.write(blob)
        with zipfile.ZipFile(zp) as z:
            names = z.namelist()
            for n in names:
                parts = n.replace('\\', '/').split('/')
                if n.startswith(('/', '\\')) or '..' in parts or ':' in parts[0]:
                    raise ApiError('File cập nhật có đường dẫn không an toàn, đã dừng')
            mains = [n for n in names if n.replace('\\', '/').split('/')[-1] == 'KhoaVeyon.py']
            if not mains:
                raise ApiError('File cập nhật không phải bản KhoaVeyon (thiếu KhoaVeyon.py)')
            root = min(mains, key=len)[:-len('KhoaVeyon.py')]
            z.extractall(os.path.join(tmp, 'x'))
        base = os.path.join(tmp, 'x', root.replace('/', os.sep))
        src_py = os.path.join(base, 'KhoaVeyon.py')
        try:
            compile(open(src_py, 'r', encoding='utf-8').read(), src_py, 'exec')
        except SyntaxError as e:
            raise ApiError('Bản cập nhật bị lỗi cú pháp, giữ nguyên bản hiện tại: %s' % e)
        todo = []
        for f in UPDATE_FILES:
            if os.path.isfile(os.path.join(base, f)):
                todo.append(f)
        for d in UPDATE_DIRS:
            for dp, _dn, fn in os.walk(os.path.join(base, d)):
                for f in fn:
                    todo.append(os.path.relpath(os.path.join(dp, f), base))
        bk = os.path.join(app_dir, 'data', 'backup-' + VERSION)
        for rel in todo:
            dst = os.path.join(app_dir, rel)
            if os.path.isfile(dst):
                os.makedirs(os.path.dirname(os.path.join(bk, rel)), exist_ok=True)
                shutil.copy2(dst, os.path.join(bk, rel))
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            shutil.copy2(os.path.join(base, rel), dst + '.new')
            os.replace(dst + '.new', dst)
        return todo
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def restart_self():
    """Mở bản mới (đợi bản này tắt xong) rồi tắt máy chủ hiện tại."""
    flags = (0x00000008 | 0x00000200) if os.name == 'nt' else 0   # DETACHED_PROCESS | CREATE_NEW_PROCESS_GROUP
    subprocess.Popen([sys.executable, os.path.abspath(__file__), '--after-update'], cwd=APP_DIR, close_fds=True,
                     creationflags=flags)
    srv = STATE.get('server')
    if srv:
        threading.Timer(0.8, srv.shutdown).start()


def startup_update_check():
    with CFG_LOCK:
        repo, auto = CFG.get('update_repo', ''), CFG.get('update_auto')
    try:
        info = check_update(repo)
    except Exception as e:  # noqa
        log('Kiểm tra bản mới lỗi: %s' % e, 'warn')
        return
    STATE['update'] = info
    if not info['newer']:
        return
    log('Có bản mới %s (đang dùng %s).' % (info['latest'], info['current']), 'warn')
    if auto and info['url']:
        try:
            files = apply_update(info['url'])
            log('Đã tự cập nhật %d file lên %s, đang khởi động lại...' % (len(files), info['latest']), 'ok')
            restart_self()
        except Exception as e:  # noqa
            log('Tự cập nhật lỗi, giữ nguyên bản cũ: %s' % e, 'err')


# ---------------------------------------------------------------- API
def clean_hosts(value):
    out = []
    for h in value or []:
        h = str(h).strip()
        if h and h not in out:
            out.append(h)
    return out


def warnings_for(kind, rows):
    res = []
    for i, r in enumerate(rows):
        a = str(r[0]).strip()
        if a:
            w = check_entry(kind, a)
            if w:
                res.append({'row': i, 'msg': w})
    return res


def list_path(kind):
    with CFG_LOCK:
        return os.path.join(abs_path(CFG['list_dir']), LISTS[kind]['file'])


def api_state(_q=None):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    cli = find_veyon_cli(cfg['veyon_cli'])
    actions = []
    for key, label, switch, _d in ACTIONS:
        actions.append({'key': key, 'label': label, 'switch': switch, 'feature': cfg['features'].get(key, ''),
                        'icon': ICONS.get(key, '▶'), 'confirm': key in CONFIRM_KEYS,
                        'kind': 'lock' if key in LOCK_KEYS else ('free' if key in FREE_KEYS else 'other')})
    lists = {}
    for k, m in LISTS.items():
        lists[k] = {'file': m['file'], 'title': m['title'], 'col_a': m['col_a'], 'hint': m['hint'],
                    'apply': m['apply'], 'apply_label': ACTION_BY_KEY[m['apply']][1],
                    'presets': {n: [list(r) for r in rows] for n, rows in m['presets'].items()}}
    return {'cfg': cfg, 'config_path': CONFIG_PATH, 'app_dir': APP_DIR, 'list_dir': abs_path(cfg['list_dir']),
            'payload_dir': PAYLOAD_DIR, 'veyon_cli_path': cli, 'veyon_cli_found': os.path.isfile(cli),
            'actions': actions, 'groups': [{'title': t, 'keys': k} for t, k in GROUPS], 'lists': lists,
            'windows': os.name == 'nt', 'version': VERSION, 'update': STATE.get('update')}


def api_list_get(q):
    kind = q.get('kind')
    if kind not in LISTS:
        raise ApiError('Danh sách không tồn tại')
    path = list_path(kind)
    try:
        rows = read_list(path)
    except Exception as e:  # noqa
        raise ApiError('Không đọc được %s: %s' % (path, e))
    return {'rows': [list(r) for r in rows], 'path': path, 'exists': os.path.exists(path),
            'warnings': warnings_for(kind, rows)}


def api_check(b):
    kind = b.get('kind')
    if kind not in LISTS:
        raise ApiError('Danh sách không tồn tại')
    return {'warnings': warnings_for(kind, b.get('rows') or [])}


def api_list_save(b):
    kind = b.get('kind')
    if kind not in LISTS:
        raise ApiError('Danh sách không tồn tại')
    rows = [(str(r[0]).strip(), str(r[1] if len(r) > 1 else '').strip()) for r in (b.get('rows') or [])
            if str(r[0]).strip()]
    path = list_path(kind)
    try:
        write_list(path, kind, rows)
    except PermissionError:
        raise ApiError('Không ghi được %s. Nếu file đang mở trong Excel, hãy đóng lại rồi lưu.' % path)
    except Exception as e:  # noqa
        raise ApiError('Không ghi được %s: %s' % (path, e))
    log('Đã lưu %d dòng vào %s' % (len(rows), path), 'ok')
    with CFG_LOCK:
        if CFG.get('auto_sync') and (CFG.get('sync_dir') or '').strip():
            try:
                n, err = sync_lists(CFG)
            except Exception as e:  # noqa
                n, err = 0, str(e)
            log('Đồng bộ lỗi: %s' % err if err else 'Đã đồng bộ %d file sang %s' % (n, CFG['sync_dir']),
                'err' if err else 'ok')
        push = bool(CFG.get('auto_push', True))
        target = (CFG.get('dist_target') or '').strip()
    hosts = clean_hosts(b.get('hosts'))
    apply_key = LISTS[kind]['apply'] if b.get('apply') else ''
    pushed = 0
    if push and hosts:
        try:
            unc_target('x', target)
        except ValueError as e:
            log('Lỗi thư mục đích, chưa gửi file xuống máy học sinh: %s' % e, 'err')
            push = False
    if push and hosts:
        pushed = len(hosts)
    elif push and not hosts:
        log('Chưa có máy học sinh nào nên chưa gửi %s xuống máy học sinh.' % os.path.basename(path), 'warn')
    if hosts and (pushed or apply_key):
        start_job(job_after_save, hosts, path, bool(pushed), apply_key, target)
    return {'ok': True, 'count': len(rows), 'warnings': warnings_for(kind, rows), 'pushed': pushed,
            'applied': bool(apply_key and hosts)}


def api_hosts(b):
    with CFG_LOCK:
        CFG['hosts'] = clean_hosts(b.get('hosts'))
        save_config(CFG)
    return {'ok': True}


def api_run(b):
    key = b.get('key')
    if key not in ACTION_BY_KEY:
        raise ApiError('Lệnh không tồn tại')
    hosts = clean_hosts(b.get('hosts'))
    if not hosts:
        raise ApiError('Chưa có máy nào để chạy')
    with CFG_LOCK:
        shot = bool(CFG.get('shot_before_lock'))
    if shot and key in LOCK_KEYS:
        start_job(job_shot_then_action, key, hosts)
    else:
        start_job(job_action, key, hosts)
    return {'ok': True}


def api_settings(b):
    with CFG_LOCK:
        new = json.loads(json.dumps(CFG))
        for k, v in b.items():
            if k in STR_KEYS:
                v = str(v).strip()
                if k == 'update_repo' and v and not re.fullmatch(r'[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+', v):
                    raise ApiError('Kho GitHub phải có dạng chủ-sở-hữu/tên-kho, ví dụ khoa/KhoaVeyon')
                if k == 'dist_mode' and v not in ('none', 'schtasks', 'veyon'):
                    raise ApiError('Cách chạy không hợp lệ')
                if k == 'dist_action' and v not in ACTION_BY_KEY:
                    raise ApiError('Mục Veyon không hợp lệ')
                new[k] = v
            elif k in BOOL_KEYS:
                new[k] = bool(v)
            elif k in LIST_KEYS:
                new[k] = [str(x) for x in (v or [])]
            elif k == 'features' and isinstance(v, dict):
                for fk, fv in v.items():
                    if fk in ACTION_BY_KEY and str(fv).strip():
                        new['features'][fk] = str(fv).strip()
        for tk_ in ('run_template', 'list_template'):
            try:
                new[tk_].format(cli='x', host='h', feature='f')
            except Exception as e:  # noqa
                raise ApiError('Mẫu lệnh "%s" không hợp lệ (chỉ dùng {cli} {host} {feature}): %s' % (tk_, e))
        CFG.clear()
        CFG.update(new)
        save_config(CFG)
    return {'ok': True}


def resolve_dist_files(b):
    files = []
    if b.get('include_lists'):
        for k in LISTS:
            p = list_path(k)
            if os.path.exists(p):
                files.append(p)
    for name in b.get('payload') or []:
        p = os.path.join(PAYLOAD_DIR, os.path.basename(str(name)))
        files.append(p)
    for p in b.get('extra') or []:
        p = str(p).strip().strip('"')
        if p:
            files.append(p)
    return files


def api_dist(b):
    do_copy = bool(b.get('copy'))
    mode = b.get('mode', 'none') if b.get('run') else 'none'
    if mode not in ('none', 'schtasks', 'veyon'):
        raise ApiError('Cách chạy không hợp lệ')
    hosts = clean_hosts(b.get('hosts'))
    if not hosts:
        raise ApiError('Chưa có máy nào để chạy')
    target = str(b.get('target') or '').strip()
    files = []
    if do_copy:
        files = resolve_dist_files(b)
        if not files:
            raise ApiError('Chưa có file nào để phân phối')
        missing = [f for f in files if not os.path.exists(abs_path(f))]
        if missing:
            raise ApiError('Không tìm thấy file:\n' + '\n'.join(missing[:8]))
        try:
            unc_target('x', target)
        except ValueError as e:
            raise ApiError(str(e))
    cmd, feature = '', ''
    if mode == 'schtasks':
        cmd = str(b.get('cmd') or '').strip()
        if not cmd:
            raise ApiError('Hãy nhập lệnh cần chạy trên máy học sinh')
    elif mode == 'veyon':
        action = b.get('action')
        if action not in ACTION_BY_KEY:
            raise ApiError('Mục Veyon không hợp lệ')
        feature = veyon_feature(action)
    if not do_copy and mode == 'none':
        raise ApiError('Chưa chọn việc cần làm')
    parts = []
    if do_copy:
        parts.append('chép %d mục sang %s' % (len(files), target))
    if mode == 'schtasks':
        parts.append('chạy lệnh: %s' % cmd)
    elif mode == 'veyon':
        parts.append('kích hoạt mục Veyon "%s"' % feature)
    start_job(job_dist, hosts, files, mode, cmd, feature, bool(b.get('interactive')), target, 'Phân phối & chạy: ' + '; '.join(parts))
    return {'ok': True}


def api_payload(_q=None):
    os.makedirs(PAYLOAD_DIR, exist_ok=True)
    items = []
    for name in sorted(os.listdir(PAYLOAD_DIR), key=str.lower):
        p = os.path.join(PAYLOAD_DIR, name)
        isdir = os.path.isdir(p)
        size = 0 if isdir else os.path.getsize(p)
        items.append({'name': name, 'dir': isdir, 'size': size})
    return {'items': items, 'dir': PAYLOAD_DIR}


def api_payload_delete(b):
    name = os.path.basename(str(b.get('name') or ''))
    p = os.path.join(PAYLOAD_DIR, name)
    if not name or not os.path.exists(p):
        raise ApiError('Không tìm thấy file')
    if os.path.isdir(p):
        shutil.rmtree(p)
    else:
        os.remove(p)
    return {'ok': True}


def api_open(b):
    with CFG_LOCK:
        folders = {'payload': PAYLOAD_DIR, 'lists': abs_path(CFG['list_dir']), 'data': os.path.dirname(CONFIG_PATH), 'shots': shot_root(), 'evidence': evidence_root()}
    p = folders.get(b.get('what'))
    if not p:
        raise ApiError('Thư mục không hợp lệ')
    os.makedirs(p, exist_ok=True)
    if os.name == 'nt':
        os.startfile(p)  # noqa
    else:
        raise ApiError('Chỉ mở được thư mục trên Windows. Đường dẫn: %s' % p)
    return {'ok': True}


def api_veyon_list(b):
    host = str(b.get('host') or '').strip()
    if not host:
        raise ApiError('Hãy thêm ít nhất một máy ở trang Điều khiển lớp học')

    def work():
        try:
            with CFG_LOCK:
                cmd = veyon_command(CFG, host, '', 'list_template')
        except Exception as e:  # noqa
            log('Lỗi mẫu lệnh: %s' % e, 'err')
            return
        rc, out = run_shell(cmd)
        log('>>> veyon-cli feature list (%s) -> mã %s\n%s' % (host, rc, out), 'head' if rc == 0 else 'err')
    start_job(work)
    return {'ok': True}


def api_setup_status(_q=None):
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    cli = find_veyon_cli(cfg['veyon_cli'])
    have, total, err = veyon_items_status(cfg)
    lists_ok = sum(1 for k in LISTS if os.path.exists(list_path(k)))
    sname = os.path.basename((cfg.get('script_path') or '').replace('\\', '/')) or 'KhoaVeyonServices.bat'
    return {'cli': os.path.isfile(cli), 'cli_path': cli, 'items_have': have, 'items_total': total, 'items_err': err,
            'admin': is_admin(), 'lists_ok': lists_ok, 'lists_total': len(LISTS),
            'script_name': sname, 'script_in_payload': os.path.isfile(os.path.join(PAYLOAD_DIR, sname)),
            'script_found': bool(find_script_file(cfg)),
            'student_bat': os.path.isfile(os.path.join(PAYLOAD_DIR, STUDENT_SETUP_BAT)),
            'hosts': len(cfg.get('hosts') or []), 'done': bool(cfg.get('setup_done')), 'windows': os.name == 'nt'}


def api_setup_run(b):
    step = b.get('step')
    hosts = clean_hosts(b.get('hosts'))
    with CFG_LOCK:
        cfg = json.loads(json.dumps(CFG))
    if step == 'all':
        start_job(job_setup_all, hosts)
    elif step == 'lists':
        n = seed_default_lists(cfg)
        log('Danh sách Excel: đã tạo %d file mẫu' % n if n else 'Danh sách Excel: đã có đủ, giữ nguyên', 'ok')
    elif step == 'payload':
        got, missing = ensure_payload(cfg)
        log('Đã tạo payload\\%s' % STUDENT_SETUP_BAT + (', lấy %s' % ', '.join(got) if got else ''), 'ok')
        if missing:
            log('Chưa tìm thấy %s: hãy thêm vào payload (trang Phân phối & chạy)' % ', '.join(missing), 'warn')
    elif step == 'veyon':
        def work():
            ok, msg = elevated_apply_veyon_items(cfg)
            log('Mục Veyon: ' + msg, 'ok' if ok else 'err')
        start_job(work)
    elif step in ('lock', 'unlock'):
        if not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        start_job(job_cmd_lock, hosts, step == 'lock')
    elif step in ('diagnose', 'push'):
        if not hosts:
            raise ApiError('Chưa có máy học sinh nào')
        start_job(job_setup_diagnose, hosts, step == 'push')
    else:
        raise ApiError('Bước không hợp lệ')
    return {'ok': True}


def api_setup_scan(_b=None):
    names, err = scan_veyon_hosts()
    if err:
        raise ApiError(err)
    added = 0
    with CFG_LOCK:
        for n in names:
            if n not in CFG['hosts']:
                CFG['hosts'].append(n)
                added += 1
        save_config(CFG)
        hosts = list(CFG['hosts'])
    log('Quét mạng: thấy %d máy chạy Veyon, thêm mới %d' % (len(names), added), 'ok' if names else 'warn')
    return {'found': names, 'added': added, 'hosts': hosts}


def api_setup_done(b):
    with CFG_LOCK:
        CFG['setup_done'] = bool(b.get('done', True))
        save_config(CFG)
    return {'ok': True}


def api_update_check(_q=None):
    with CFG_LOCK:
        repo = CFG.get('update_repo', '')
    info = check_update(repo)
    STATE['update'] = info
    return info


def api_update_apply(_b=None):
    info = STATE.get('update') or {}
    if not info.get('newer') or not info.get('url'):
        raise ApiError('Chưa có bản mới để cập nhật (hãy bấm kiểm tra trước)')

    def work():
        try:
            files = apply_update(info['url'])
            log('Đã cập nhật %d file lên %s, đang khởi động lại...' % (len(files), info['latest']), 'ok')
            restart_self()
        except Exception as e:  # noqa
            log('Cập nhật lỗi, giữ nguyên bản cũ: %s' % e, 'err')
    start_job(work)
    return {'ok': True}


def api_sync(_b=None):
    with CFG_LOCK:
        try:
            n, err = sync_lists(CFG)
        except Exception as e:  # noqa
            n, err = 0, str(e)
        dst = CFG.get('sync_dir')
    if err:
        raise ApiError(err)
    log('Đã đồng bộ %d file sang %s' % (n, dst), 'ok')
    return {'ok': True, 'count': n}


def api_log(q):
    STATE['last_ping'] = time.time()
    STATE['pinged'] = True
    since = int(q.get('since') or 0)
    with LOG_LOCK:
        items = [e for e in LOG if e['id'] > since]
        last = LOG_SEQ[0]
    return {'items': items, 'last': last, 'busy': STATE['busy']}


def api_ping(_q=None):
    STATE['last_ping'] = time.time()
    STATE['pinged'] = True
    return {'ok': True}


def api_quit(_b=None):
    srv = STATE.get('server')
    if srv:
        threading.Timer(0.4, srv.shutdown).start()
    return {'ok': True}


GET_ROUTES = {'/api/update_check': api_update_check, '/api/setup_status': api_setup_status, '/api/state': api_state, '/api/list': api_list_get, '/api/payload': api_payload,
              '/api/log': api_log, '/api/ping': api_ping, '/api/vcfg_info': api_vcfg_info}
POST_ROUTES = {'/api/update_apply': api_update_apply, '/api/setup_run': api_setup_run, '/api/setup_scan': api_setup_scan, '/api/setup_done': api_setup_done,
               '/api/list': api_list_save, '/api/check': api_check, '/api/hosts': api_hosts, '/api/run': api_run,
               '/api/settings': api_settings, '/api/dist': api_dist, '/api/open': api_open,
               '/api/veyon_list': api_veyon_list, '/api/sync': api_sync, '/api/quit': api_quit,
               '/api/payload_delete': api_payload_delete, '/api/vcfg_save': api_vcfg_save, '/api/vcfg_run': api_vcfg_run, '/api/shot': api_shot, '/api/watch': api_watch}


class Handler(http.server.BaseHTTPRequestHandler):
    server_version = 'KhoaVeyon'
    protocol_version = 'HTTP/1.1'

    def log_message(self, *args):  # không ghi log truy cập
        pass

    def _send(self, code, body, ctype='application/json; charset=utf-8'):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False).encode('utf-8')
        elif isinstance(body, str):
            body = body.encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        if self.command != 'HEAD':
            self.wfile.write(body)

    def _host_ok(self):
        h = (self.headers.get('Host') or '').lower()
        return h in ('127.0.0.1:%d' % STATE['port'], 'localhost:%d' % STATE['port'])

    def _token_ok(self):
        return secrets.compare_digest(self.headers.get('X-Token') or '', STATE['token'])

    def _guard(self):
        if not self._host_ok():
            self._send(403, {'error': 'Host không hợp lệ'})
            return False
        return True

    def _api(self, fn, arg):
        try:
            self._send(200, fn(arg))
        except ApiError as e:
            self._send(400, {'error': str(e)})
        except Exception as e:  # noqa
            self._send(500, {'error': '%s: %s' % (type(e).__name__, e)})

    def do_GET(self):
        if not self._guard():
            return
        u = urllib.parse.urlparse(self.path)
        q = dict(urllib.parse.parse_qsl(u.query))
        if u.path == '/':
            if not secrets.compare_digest(q.get('t', ''), STATE['token']):
                self._send(403, '<meta charset="utf-8"><h3>Hãy mở chương trình bằng Chay-KhoaVeyon.bat</h3>',
                           'text/html; charset=utf-8')
                return
            try:
                with open(os.path.join(WEB_DIR, 'index.html'), 'r', encoding='utf-8') as f:
                    html = f.read()
            except OSError:
                self._send(500, '<meta charset="utf-8"><h3>Thiếu thư mục web (web\\index.html)</h3>',
                           'text/html; charset=utf-8')
                return
            self._send(200, html.replace('__TOKEN__', STATE['token']), 'text/html; charset=utf-8')
        elif u.path == '/favicon.ico':
            self._send(204, b'', 'image/x-icon')
        elif u.path in GET_ROUTES:
            if not self._token_ok():
                self._send(403, {'error': 'Sai mã bảo vệ'})
                return
            self._api(GET_ROUTES[u.path], q)
        else:
            self._send(404, {'error': 'Không tìm thấy'})

    def do_POST(self):
        if not self._guard():
            return
        u = urllib.parse.urlparse(self.path)
        n = int(self.headers.get('Content-Length') or 0)
        raw = self.rfile.read(n) if n else b''
        if u.path not in POST_ROUTES:
            self._send(404, {'error': 'Không tìm thấy'})
            return
        if not self._token_ok():
            self._send(403, {'error': 'Sai mã bảo vệ'})
            return
        try:
            body = json.loads(raw.decode('utf-8') or '{}')
        except Exception:  # noqa
            self._send(400, {'error': 'Dữ liệu gửi lên không hợp lệ'})
            return
        self._api(POST_ROUTES[u.path], body if isinstance(body, dict) else {})

    def do_PUT(self):
        """Tải file lên thư mục payload: PUT /api/upload?name=... (nội dung file là phần thân)."""
        if not self._guard():
            return
        u = urllib.parse.urlparse(self.path)
        n = int(self.headers.get('Content-Length') or 0)
        if u.path != '/api/upload' or not self._token_ok():
            self.rfile.read(min(n, MAX_UPLOAD))
            self._send(403, {'error': 'Không được phép'})
            return
        if n > MAX_UPLOAD:
            self._send(413, {'error': 'File quá lớn (tối đa 500 MB)'})
            return
        name = os.path.basename(urllib.parse.unquote(dict(urllib.parse.parse_qsl(u.query)).get('name', '')))
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .')
        if not name:
            self.rfile.read(n)
            self._send(400, {'error': 'Thiếu tên file'})
            return
        os.makedirs(PAYLOAD_DIR, exist_ok=True)
        dest = os.path.join(PAYLOAD_DIR, name)
        left = n
        with open(dest, 'wb') as f:
            while left > 0:
                chunk = self.rfile.read(min(1024 * 1024, left))
                if not chunk:
                    break
                f.write(chunk)
                left -= len(chunk)
        log('Đã thêm file vào thư mục payload: %s' % name, 'ok')
        self._send(200, {'ok': True, 'name': name})


class Server(http.server.ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = False


def make_server():
    for port in (8765, 0):
        try:
            return Server(('127.0.0.1', port), Handler)
        except OSError:
            continue
    raise RuntimeError('Không mở được cổng mạng nội bộ')


def _info_path():
    return os.path.join(os.path.dirname(CONFIG_PATH), 'server.json')


def existing_instance():
    """Nếu KhoaVeyon đã chạy sẵn thì trả về địa chỉ của nó để mở lại cửa sổ thay vì chạy bản thứ hai."""
    try:
        with open(_info_path(), 'r', encoding='utf-8') as f:
            info = json.load(f)
        req = urllib.request.Request('http://127.0.0.1:%d/api/ping' % info['port'], headers={'X-Token': info['token']})
        with urllib.request.urlopen(req, timeout=1.5) as r:
            if r.status == 200:
                return 'http://127.0.0.1:%d/?t=%s' % (info['port'], info['token'])
    except Exception:  # noqa
        pass
    return ''


def browser_exe():
    for env in ('ProgramFiles(x86)', 'ProgramFiles', 'LOCALAPPDATA'):
        base = os.environ.get(env)
        if not base:
            continue
        for rel in (('Microsoft', 'Edge', 'Application', 'msedge.exe'),
                    ('Google', 'Chrome', 'Application', 'chrome.exe')):
            p = os.path.join(base, *rel)
            if os.path.isfile(p):
                return p
    return ''


def open_ui(url):
    """Mở như một cửa sổ ứng dụng (Edge/Chrome --app, không có thanh địa chỉ); không có thì mở trình duyệt mặc định."""
    exe = browser_exe() if os.name == 'nt' else ''
    if exe:
        try:
            subprocess.Popen([exe, '--app=' + url, '--window-size=1360,900'])
            return
        except Exception:  # noqa
            pass
    webbrowser.open(url)


def watchdog(srv):
    while True:
        time.sleep(5)
        if STATE['pinged'] and STATE['busy'] == 0 and time.time() - STATE['last_ping'] > IDLE_EXIT_SECONDS:
            srv.shutdown()
            return


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    if '--after-update' in argv:  # bản cũ đang tắt, đợi nó nhả cổng rồi chạy bản mới
        for _ in range(40):
            if not existing_instance():
                break
            time.sleep(0.5)
    if '--apply-veyon-items' in argv:
        run_apply_mode()
        return
    if '--elevated-job' in argv:
        run_elevated_job()
        return
    os.makedirs(PAYLOAD_DIR, exist_ok=True)
    inst = existing_instance()
    if inst:
        if '--no-browser' not in argv:
            open_ui(inst)
        return
    srv = make_server()
    STATE['server'] = srv
    STATE['port'] = srv.server_address[1]
    STATE['token'] = secrets.token_urlsafe(18)
    url = 'http://127.0.0.1:%d/?t=%s' % (STATE['port'], STATE['token'])
    try:
        with open(_info_path(), 'w', encoding='utf-8') as f:
            json.dump({'port': STATE['port'], 'token': STATE['token']}, f)
    except OSError:
        pass
    log('Sẵn sàng. Thư mục danh sách: %s' % abs_path(CFG['list_dir']), 'head')
    log('Cài đặt lưu tại: %s' % CONFIG_PATH)
    threading.Thread(target=watchdog, args=(srv,), daemon=True).start()
    if CFG.get('update_check') and (CFG.get('update_repo') or '').strip():
        threading.Thread(target=startup_update_check, daemon=True).start()
    if '--no-browser' in argv:
        if sys.stdout:
            print(url, flush=True)
    else:
        open_ui(url)
    try:
        srv.serve_forever()
    finally:
        try:
            os.remove(_info_path())
        except OSError:
            pass


if __name__ == '__main__':
    try:
        main()
    except Exception:  # noqa
        try:
            with open(os.path.join(os.path.dirname(CONFIG_PATH), 'error.log'), 'a', encoding='utf-8') as f:
                f.write(time.strftime('%Y-%m-%d %H:%M:%S') + '\n' + traceback.format_exc() + '\n')
        except OSError:
            pass
        raise
